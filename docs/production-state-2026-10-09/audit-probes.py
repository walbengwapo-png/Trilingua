"""Offline reproductions. Run against the disposable application copy."""
from pathlib import Path
import ast, json, os, sys, tempfile
from unittest.mock import Mock, patch

sys.stdout.reconfigure(encoding='utf-8')

app = Path(sys.argv[1]).resolve()
sys.path.insert(0, str(app / 'Model'))
os.environ['TRANSLATION_CACHE_ENABLED'] = 'false'
os.environ['GPTOSS_MAX_ATTEMPTS'] = '3'
from cache.sqlite_cache import SQLiteTranslationCache
from memory.translation_cache import TranslationCache
from document.reconstructor import translate_xlsx_inplace
from pipeline.translation_pipeline import TranslationPipeline
from validators.ai_quality_reviewer import AIQualityReviewer
from dto.responses import TranslationResponse
from providers.gptoss import GPTOSSProvider
from providers.fallback import FallbackTranslationProvider
from openpyxl import Workbook, load_workbook
import requests

results = {}
with tempfile.TemporaryDirectory() as scratch:
    source, output = Path(scratch) / 'in.xlsx', Path(scratch) / 'out.xlsx'
    workbook = Workbook()
    workbook.active['A1'] = '=SUM(1,2)'
    workbook.active['B1'] = 'Hello'
    workbook.save(source)
    translate_xlsx_inplace(str(source), str(output), lambda text, **kw: 'Kumusta')
    output_book = load_workbook(output, data_only=False)
    results['xlsx_formula_loss'] = {'source_A1': '=SUM(1,2)', 'output_A1': output_book.active['A1'].value, 'output_B1': output_book.active['B1'].value}
    assert output_book.active['A1'].value is None
    output_book.close()
    cache = SQLiteTranslationCache(db_path=str(Path(scratch) / 'cache.db'))
    cache.put('US', 'Cebuano', 'mock', 'Estados Unidos', source_lang='English')
    cache.flush()
    results['case_collision'] = {'requested': 'us', 'returned': cache.get('us', 'Cebuano', 'mock', source_lang='English'), 'memory_keys_equal': TranslationCache._make_key('US', 'English', 'Cebuano') == TranslationCache._make_key('us', 'English', 'Cebuano')}
    assert results['case_collision']['returned'] == 'Estados Unidos'

class FakeProvider:
    name, model_name, max_concurrency = 'mock', 'mock-model', 1
    batch_limits = {'max_batch_chars': 1500, 'max_batch_items': 5}
    def __init__(self): self.calls = 0
    def translate(self, text, **kw):
        self.calls += 1
        return TranslationResponse(translated_text=text, provider=self.name, model=self.model_name)
    def estimate_tokens(self, text): return len(text.split())

provider = FakeProvider()
pipeline = TranslationPipeline(provider)
out = pipeline.batch_translate_blocks([{'text':'I','type':'paragraph','block_index':0}], 'English', 'Cebuano')
results['single_character_document_passthrough'] = {'output': out[0]['text'], 'provider_calls': provider.calls}
assert out[0]['text'] == 'I' and provider.calls == 0
echo = pipeline._translate_with_echo_guard(text='The dog is hungry.', source_lang='English', target_lang='Cebuano', block_type='paragraph', context_hint='', document_type='')
results['persistent_echo_success'] = {'text': echo.translated_text, 'success': echo.success, 'provider_calls': provider.calls, 'warnings': echo.warnings}
assert echo.success and echo.translated_text == 'The dog is hungry.'
review = AIQualityReviewer(Mock()).review('tubig', 'tubig')
results['identity_review'] = {'score':review.score,'retry_required':review.retry_required,'issues':[i.category for i in review.issues]}
assert review.retry_required and 'untranslated' in results['identity_review']['issues']

results['terminal_http_retry'] = {}
for status in (401, 403, 429):
    response = requests.Response()
    response.status_code = status
    response.url = 'http://127.0.0.1:1/api/chat'
    response._content = b'{"error":"offline fixture"}'
    current = GPTOSSProvider(api_url=response.url)
    with patch.object(current._session, 'post', return_value=response) as post, patch('providers.gptoss._time.sleep'):
        failure = current.translate('Hello there.', 'English', 'Cebuano')
    results['terminal_http_retry'][str(status)] = {'attempts':post.call_count,'success':failure.success}
    assert post.call_count == 3 and not failure.success

primary, backup = FakeProvider(), FakeProvider()
primary.name, backup.name = 'gptoss', 'gemini'
primary.translate = Mock(return_value=TranslationResponse(translated_text='',provider='gptoss',success=False,error_message='GPT-OSS rate limit reached after paced retries'))
backup.translate = Mock(return_value=TranslationResponse(translated_text='Kumusta',provider='gemini'))
fallback = FallbackTranslationProvider(primary, backup)
switched = fallback.translate('Hello', 'English', 'Cebuano')
results['quota_fallback'] = {'provider':switched.provider,'backup_calls':backup.translate.call_count}
assert switched.provider == 'gemini' and backup.translate.call_count == 1

current = GPTOSSProvider(api_url='http://127.0.0.1:1/api/chat')
response = Mock(status_code=200)
response.json.return_value = {'message':{'content':'Kumusta ka.'}}
with patch.object(current._session,'post',return_value=response):
    translated = current.translate('Hello there.', 'English', 'Cebuano')
results['unreported_tokens_estimated'] = translated.token_usage
assert translated.token_usage == {'input':2,'output':2}

tree = ast.parse((app / 'Model/server.py').read_text(encoding='utf-8'))
results['startup_warmup'] = {'top_level_call_lines':[n.lineno for n in tree.body if isinstance(n,ast.Expr) and isinstance(n.value,ast.Call) and isinstance(n.value.func,ast.Name) and n.value.func.id == '_warm_cold_start']}
assert results['startup_warmup']['top_level_call_lines'] == [307]
result_path = Path(__file__).with_name('python-probes.json')
result_path.write_text(json.dumps(results,indent=2),encoding='utf-8')
print(json.dumps(results,indent=2))

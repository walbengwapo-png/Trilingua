"""Build the consolidated current-state TriLingua evaluation report."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\dev\Trilingua\chatgpt\test")
OUT_DOCX = ROOT / "Trilingua Current Evaluation Report.docx"
OUT_MD = ROOT / "comparison" / "Trilingua-Current-Evaluation-Report.md"
BLUE = "17365D"
PALE_BLUE = "EAF2F8"
GRID = "D9D9D9"


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def minutes(value: float | None) -> str:
    if value is None:
        return "Not run"
    return f"{value / 60000:.2f} min"


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    fill_element = OxmlElement("w:shd")
    fill_element.set(qn("w:fill"), fill)
    properties.append(fill_element)


def borders(table) -> None:
    table_properties = table._tbl.tblPr
    table_borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), GRID)
        table_borders.append(item)
    table_properties.append(table_borders)


def set_cell(cell, text: str, *, header: bool = False) -> None:
    cell.text = text
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.space_before = Pt(2)
    run = paragraph.runs[0]
    run.font.name = "Aptos"
    run.font.size = Pt(8.5)
    if header:
        shade(cell, BLUE)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
    else:
        run.font.color.rgb = RGBColor(0, 0, 0)


def add_table(doc: Document, headings: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headings))
    table.autofit = True
    borders(table)
    for i, heading in enumerate(headings):
        set_cell(table.rows[0].cells[i], heading, header=True)
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell(cells[i], value)
            if index % 2:
                shade(cells[i], PALE_BLUE)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def add_heading(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Heading 1")
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    run.font.name = "Aptos Display"
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_body(doc: Document, text: str, bold_start: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.12
    if bold_start and text.startswith(bold_start):
        paragraph.add_run(bold_start).bold = True
        paragraph.add_run(text[len(bold_start):])
    else:
        paragraph.add_run(text)
    for run in paragraph.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(10.5)


def find_run(data: dict, source: str, target: str) -> dict | None:
    return next(
        (row for row in data.get("full_runs", [])
         if row.get("source") == source and row.get("target_language") == target),
        None,
    )


def single_metadata(language_dir: str) -> dict | None:
    path = ROOT / "sona-four-page-excerpt" / "gptoss" / language_dir / "translation-metadata.json"
    return load_json(path) if path.is_file() else None


def main() -> None:
    nllb = load_json(ROOT / "iteration-1-nllb-200-600m" / "findings.json")
    gpt = load_json(ROOT / "iteration-2-gptoss" / "findings.json")
    dfs = load_json(ROOT / "ad-hoc-gptoss-dfs" / "translation-metadata.json")
    excerpt_fil = single_metadata("filipino-content")
    excerpt_ceb = single_metadata("cebuano-content")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title.add_run("TriLingua Current Translation Evaluation Report")
    title_run.font.name = "Aptos Display"
    title_run.font.color.rgb = RGBColor(0, 0, 0)
    subtitle = doc.add_paragraph("Iteration 1 NLLB baseline, Iteration 2 GPT OSS, DFS document check, and reduced SONA feasibility test")
    subtitle.paragraph_format.space_after = Pt(10)
    for run in subtitle.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(10)
        run.font.italic = True

    add_body(
        doc,
        "Decision: keep Iteration 3 paused. NLLB is the only provider that completed the full SONA test in both target languages. GPT OSS has not produced a valid SONA output for comparison because its unchanged-source safety guard rejects title and header blocks, and its full-document run is substantially slower.",
        "Decision:",
    )

    add_heading(doc, "Scope and Evaluation Method")
    add_body(doc, "The evaluation covers English to Filipino and English to Cebuano only. Reverse translation directions are not part of this run. Each provider was selected directly with translation fallback and cache disabled. SONA quality scores compare normalized extracted text against the supplied official Filipino and Cebuano editions. Structural validation checks page count and non-empty extracted text; it is not a human linguistic or visual review.")
    add_body(doc, "The time-boxed SONA test uses physical PDF pages 3 through 6: the first four speech-content pages. Cover and front-matter pages were excluded because titles, proper names, and branding can legitimately remain unchanged and therefore distort the unchanged-source safety check.")

    add_heading(doc, "Full Iteration Results")
    nllb_fil = find_run(nllb, "SONA_2019_English.pdf", "Filipino")
    nllb_ceb = find_run(nllb, "SONA_2019_English.pdf", "Cebuano")
    gpt_fil = find_run(gpt, "SONA_2019_English.pdf", "Filipino")
    add_table(doc, ["Provider", "SONA target", "Result", "Time", "BLEU", "chrF", "Structure"], [
        ["NLLB 200 600M", "Filipino", "Completed", minutes(nllb_fil.get("elapsed_ms")), str(nllb_fil["quality"]["bleu"]), str(nllb_fil["quality"]["chrf"]), "Pass"],
        ["NLLB 200 600M", "Cebuano", "Completed", minutes(nllb_ceb.get("elapsed_ms")), str(nllb_ceb["quality"]["bleu"]), str(nllb_ceb["quality"]["chrf"]), "Pass"],
        ["GPT OSS 20B cloud", "Filipino", "Rejected", minutes(gpt_fil.get("elapsed_ms")), "Not available", "Not available", "No PDF"],
        ["GPT OSS 20B cloud", "Cebuano", "Not run", "Not run", "Not available", "Not available", "Not run"],
    ])
    add_body(doc, "NLLB used facebook/nllb-200-distilled-600M at revision f8d333a098d19b4fd9a8b18f94170487ad3f821d, running locally on CPU. GPT OSS used gpt-oss:20b-cloud through the current service. GPT OSS Filipino SONA was rejected after two blocks remained unchanged, after 89.84 minutes; no false-success document was created.")

    add_heading(doc, "Reduced Four Page SONA Feasibility Test")
    excerpt_rows = []
    for label, record in (("Filipino", excerpt_fil), ("Cebuano", excerpt_ceb)):
        if record:
            message = record.get("engine_response", {}).get("error", "")
            unchanged = "unchanged block" if "unchanged" in message else "No safety error"
            excerpt_rows.append([label, record.get("status", "Unknown").title(), minutes(record.get("elapsed_ms")), unchanged, "No PDF" if not record.get("output_path") else "PDF produced"])
        else:
            excerpt_rows.append([label, "In progress or not run", "Pending", "Pending", "Pending"])
    add_table(doc, ["Target", "Result", "Elapsed", "Safety-check outcome", "Output"], excerpt_rows)
    add_body(doc, "A preliminary physical-pages-1-to-4 probe was excluded from comparison because the cover and front matter contained expected unchanged title and branding text. It was safely rejected in 6.63 minutes for Filipino and 8.94 minutes for Cebuano. The representative speech-content sample replaces that probe for the feasibility finding.")
    add_body(doc, "The four-page sample cannot receive BLEU or chrF in its current state because the current logic returns no translated PDF when a title/header block remains unchanged. This is a functional limitation in the validation policy, not a claim that the model translated the body incorrectly.")

    add_heading(doc, "Other Completed Evidence")
    add_table(doc, ["Source", "Provider", "Target", "Result", "Time", "Structural check"], [
        ["PSF Brochure", "NLLB 200 600M", "Filipino", "Completed", "4.69 min", "Pass"],
        ["PSF Brochure", "NLLB 200 600M", "Cebuano", "Rejected", "4.92 min", "No PDF"],
        ["PSF Brochure", "GPT OSS 20B cloud", "Filipino", "Rejected", "11.47 min", "No PDF"],
        ["PSF Brochure", "GPT OSS 20B cloud", "Cebuano", "Completed", "21.76 min", "Pass"],
        ["DFS Website Foundation Agreement", "GPT OSS 20B cloud", "Cebuano", "Completed", minutes(dfs.get("elapsed_ms")), "Pass; 4 pages; non-empty text"],
    ])
    add_body(doc, "The PSF Brochure has no verified Filipino or Cebuano reference translation, so quality scores are intentionally not assigned. The DFS result confirms that GPT OSS can complete a four-page document under the isolated setup, but it does not establish SONA quality because there is no aligned reference translation for that agreement.")

    add_heading(doc, "Findings and Recommendation")
    add_body(doc, "NLLB is currently the operationally demonstrated baseline for full SONA translation. It completed both target languages, passed structural checks, and showed stable output across three two-page repeats per language. Its median two-page latency was 1.61 minutes for Filipino and 1.52 minutes for Cebuano.")
    add_body(doc, "GPT OSS is not eligible for a future translation-logic update based on the current evidence. It has no valid full SONA quality score, took 89.84 minutes before the failed Filipino full run ended, and the reduced SONA sample exposed a blocker in the unchanged-source policy. Its successful DFS and PSF Cebuano outputs demonstrate that the provider can translate some PDFs, but not that it is reliable for this SONA benchmark.")
    add_body(doc, "Before any further GPT OSS quality test, improve the unchanged-source validation so approved non-translatable spans such as document titles, named entities, and headers can pass while genuine body-text echoes still fail. Re-run the four speech-content SONA pages in both targets after that change. Only then calculate excerpt BLEU and chrF against the matching official reference pages and decide whether a full SONA rerun is justified.")

    add_heading(doc, "Artifact Locations")
    add_body(doc, "Original files remain unchanged under Original Files. Existing iteration artifacts are stored in iteration-1-nllb-200-600m and iteration-2-gptoss. The reduced SONA inputs and metadata are in sona-four-page-excerpt. The standalone GPT OSS DFS Cebuano output and its report are in ad-hoc-gptoss-dfs and the test root.")
    doc.save(OUT_DOCX)

    markdown = """# TriLingua Current Translation Evaluation Report\n\n## Decision\n\nKeep Iteration 3 paused. NLLB is the only provider that completed full SONA translations in both target languages. GPT OSS has no valid SONA quality comparison because its unchanged-source safeguard rejects title and header blocks, and its full Filipino SONA run took 89.84 minutes before rejection.\n\n## Full SONA Results\n\n| Provider | Target | Result | Time | BLEU | chrF | Structure |\n|---|---|---|---:|---:|---:|---|\n| NLLB 200 600M | Filipino | Completed | 25.49 min | 29.42 | 69.47 | Pass |\n| NLLB 200 600M | Cebuano | Completed | 29.39 min | 19.08 | 65.06 | Pass |\n| GPT OSS 20B cloud | Filipino | Rejected | 89.84 min | N/A | N/A | No PDF |\n| GPT OSS 20B cloud | Cebuano | Not run | N/A | N/A | N/A | N/A |\n\n## Reduced Four Page SONA Test\n\nThe representative sample is physical PDF pages 3 to 6, the first four speech-content pages. It excludes cover/front matter that contain legitimate unchanged titles and branding.\n\n"""
    for label, record in (("Filipino", excerpt_fil), ("Cebuano", excerpt_ceb)):
        if record:
            markdown += f"- {label}: {record.get('status')} in {minutes(record.get('elapsed_ms'))}; no output PDF when the unchanged-source safety check rejects a block.\n"
        else:
            markdown += f"- {label}: in progress or not run when this report was generated.\n"
    markdown += """\n## Recommendation\n\nDo not promote GPT OSS or begin Iteration 3. First revise the unchanged-source validation to permit documented non-translatable titles, proper names, and repeating headers without accepting genuine untranslated body text. Then rerun the four-page content sample in both languages and score it against matching official reference pages before considering a full SONA rerun.\n\n## Limitations\n\nScope is English to Filipino and English to Cebuano only. BLEU and chrF use normalized extracted PDF text and do not replace bilingual human review. PSF and DFS lack verified aligned references and therefore have no quality score.\n"""
    OUT_MD.write_text(markdown, encoding="utf-8")
    print(OUT_DOCX)
    print(OUT_MD)


if __name__ == "__main__":
    main()

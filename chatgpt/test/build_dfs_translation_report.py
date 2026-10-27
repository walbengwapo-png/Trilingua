from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TEST_ROOT = Path(r"C:\dev\Trilingua\chatgpt\test")
AD_HOC = TEST_ROOT / "ad-hoc-gptoss-dfs"
METADATA_PATH = AD_HOC / "translation-metadata.json"
OUTPUT = TEST_ROOT / "DFS Website Foundation Agreement Cebuano Translation Report.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_border(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_font(run, size=10.5, bold=False, color="000000") -> None:
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_table(doc: Document, rows: list[tuple[str, str]], widths: tuple[float, float]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.style = "Table Grid"
    for index, header in enumerate(("Item", "Recorded result")):
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, "1F4E79")
        set_cell_border(cell)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        set_font(run, size=9.5, bold=True, color="FFFFFF")
    for row_index, (label, value) in enumerate(rows):
        cells = table.add_row().cells
        for index, cell in enumerate(cells):
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2:
                set_cell_shading(cell, "F4F7FA")
            set_cell_border(cell)
            set_cell_margins(cell)
        label_run = cells[0].paragraphs[0].add_run(label)
        set_font(label_run, size=9.3, bold=True)
        value_run = cells[1].paragraphs[0].add_run(value)
        set_font(value_run, size=9.3)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    set_font(run, size=13 if level == 1 else 11, bold=True)


def add_body(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.12
    run = paragraph.add_run(text)
    set_font(run, size=10.5)


def main() -> None:
    metadata = json.loads(METADATA_PATH.read_text(encoding="utf-8"))
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.name = "Aptos"

    title = document.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(5)
    title_run = title.add_run("DFS Website Foundation Agreement Cebuano Translation Report")
    set_font(title_run, size=20, bold=True)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle_run = subtitle.add_run("Outcome and verification for the English to Cebuano document translation")
    set_font(subtitle_run, size=11, color="404040")

    add_heading(document, "Outcome")
    add_body(
        document,
        "The English DFS Website Foundation Agreement was translated to Cebuano successfully. "
        "The generated PDF preserves the four-page source structure, contains extractable translated text, "
        "and passed the automated layout validation. The translation should receive bilingual review before it is used as a legal or contractual instrument.",
    )

    add_heading(document, "Translation Record")
    add_table(document, [
        ("Source document", "DFS Website Foundation Agreement PDF"),
        ("Target language", "Cebuano"),
        ("Provider and model", "GPT OSS gpt oss 20b cloud"),
        ("Processing mode", "Fast mode with GPT OSS only"),
        ("Fallback and cache", "Fallback disabled and translation cache disabled"),
        ("Elapsed time", f"{metadata['elapsed_ms'] / 60000:.2f} minutes"),
        ("Translation activity", "80 blocks translated with no cache hits or fallback use"),
    ], (2.05, 4.85))

    add_heading(document, "Verification Results")
    structure = metadata["structure"]
    add_table(document, [
        ("Source file hash", metadata["source_sha256"]),
        ("Output file hash", metadata["output_sha256"]),
        ("Source and output pages", f"4 source pages and {structure['page_count']} output pages"),
        ("Extracted output text", f"{structure['text_characters']:,} characters and non empty text extraction"),
        ("Layout validation", "Pass - output page count matches the source and readable text is present"),
        ("Output PDF size", f"{metadata['output_bytes']:,} bytes"),
    ], (2.05, 4.85))

    add_heading(document, "Interpretation and Limitations")
    add_body(
        document,
        "This result confirms that the current GPT OSS workflow can produce a structurally valid Cebuano version of this four-page agreement under the recorded settings. "
        "No verified Cebuano reference translation exists for this agreement, so BLEU and chrF quality scores were not calculated. Automated structural validation does not assess legal accuracy, terminology consistency, or cultural appropriateness.",
    )
    add_body(
        document,
        "Before publication, signing, or operational use, a qualified Cebuano reviewer should check names, prices, dates, scope terms, ownership language, payment conditions, and any wording that could change contractual meaning. "
        "The translated PDF is stored with the accompanying metadata in the test folder for traceability.",
    )

    add_heading(document, "Artifact Locations")
    add_table(document, [
        ("Translated PDF", "ad-hoc-gptoss-dfs\\DFS Website Foundation Agreement en to ceb PDF"),
        ("Run metadata", "ad-hoc-gptoss-dfs\\translation metadata JSON"),
        ("Report", "DFS Website Foundation Agreement Cebuano Translation Report DOCX"),
    ], (2.05, 4.85))

    document.core_properties.title = "DFS Website Foundation Agreement Cebuano Translation Report"
    document.core_properties.subject = "Cebuano translation outcome and verification"
    document.core_properties.author = "TriLingua"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()

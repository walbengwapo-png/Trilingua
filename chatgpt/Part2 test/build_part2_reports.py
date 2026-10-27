"""Create Markdown and DOCX reports from the Part2 evaluation evidence."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
PART2 = ROOT / "chatgpt" / "Part2 test"
ITERATION = PART2 / "iteration-1-nllb-200-600m"
BLUE = "17365D"
PALE_BLUE = "EAF2F8"
GRID = "D9D9D9"


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    properties.append(element)


def borders(table) -> None:
    props = table._tbl.tblPr
    element = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        part = OxmlElement(f"w:{edge}")
        part.set(qn("w:val"), "single")
        part.set(qn("w:sz"), "4")
        part.set(qn("w:color"), GRID)
        element.append(part)
    props.append(element)


def set_run(run, size: float = 10.0, bold: bool = False, color: str = "000000") -> None:
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_body(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.1
    set_run(paragraph.add_run(text), 10.2)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    paragraph.paragraph_format.space_after = Pt(4)
    set_run(paragraph.add_run(text), 13 if level == 1 else 11, True)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = True
    borders(table)
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade(cell, BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(paragraph.add_run(value), 8.6, True, "FFFFFF")
    for row_number, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_number % 2:
                shade(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            set_run(paragraph.add_run(value), 8.6)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def new_doc(title: str, subtitle: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.2)
    paragraph = doc.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(paragraph.add_run(title), 20, True)
    subtitle_paragraph = doc.add_paragraph()
    subtitle_paragraph.paragraph_format.space_after = Pt(12)
    set_run(subtitle_paragraph.add_run(subtitle), 10.5, False, "404040")
    return doc


def minutes(value: float | None) -> str:
    return "Not available" if value is None else f"{value / 60000:.2f} min"


def report_iteration(findings: dict, metadata: dict) -> tuple[str, Document]:
    rows = []
    for row in findings["full_runs"]:
        quality = row.get("quality", {})
        structure = row.get("structure", {})
        rows.append([
            f"{row['source_language']} to {row['target_language']}", row["status"].title(),
            minutes(row.get("elapsed_ms")), str(quality.get("bleu", "N/A")),
            str(quality.get("chrf", "N/A")), structure.get("layout_validation", "N/A").title(),
        ])
    stability = []
    for direction, result in findings["stability_runs"].items():
        stability.append([direction, f"{result['successful_runs']}/3", minutes(result.get("median_latency_ms")), str(result["distinct_output_hashes"]), str(result["stable_output"])])
    model = metadata.get("model", findings["health"].get("model", "Unavailable"))
    markdown = f"""# Part2 Iteration 1 NLLB Evaluation

## Outcome

The Part2 NLLB evaluation completed the English-to-Cebuano and Cebuano-to-English translations using the supplied 25-page parallel story-book editions. The accompanying JSON artifacts preserve source/reference hashes, engine evidence, output hashes, structural checks, quality scores, and repeat-run results.

## Configuration

- Provider: `nllb`
- Model: `{model}`
- Processing: `fast` mode, automatic PDF column detection
- Fallback: disabled
- Translation cache: disabled
- Quality method: corpus BLEU and chrF using all extracted PDF pages with whitespace normalization only

## Full Document Results

| Direction | Status | Elapsed | BLEU | chrF | Structural check |
|---|---|---:|---:|---:|---|
""" + "\n".join("| " + " | ".join(row) + " |" for row in rows) + """

## Stability Results

| Direction | Successful runs | Median latency | Output hashes | Stable output |
|---|---:|---:|---:|---|
""" + "\n".join("| " + " | ".join(row) + " |" for row in stability) + """

## Method and Limits

The English original was scored against the supplied Cebuano edition, and the Cebuano edition was scored against the supplied English original. BLEU and chrF measure overlap with those particular editions; they do not replace bilingual review for naturalness, theology, age appropriateness, or page-level visual quality. Structural validation verifies page count and extractable text, while the rendered-output review checks for visible layout defects.

## Traceability

The report is reconciled with `iteration-1-nllb-200-600m/benchmark-manifest.json`, `findings.json`, and `run-metadata.json`. The source and reference PDFs are not modified.
"""
    doc = new_doc("Part2 Iteration 1 NLLB Evaluation", "NLLB translation results for the English and Cebuano parallel story-book PDFs")
    add_body(doc, "The evaluation completed both Part2 directions with the isolated NLLB provider. The detailed metadata records reproducible configuration, document outputs, hashes, scores, and stability measurements.")
    add_heading(doc, "Configuration")
    add_table(doc, ["Item", "Recorded value"], [["Provider", "NLLB"], ["Model", str(model)], ["Processing", "Fast mode; automatic PDF columns"], ["Fallback and cache", "Both disabled"], ["Scoring", "Corpus BLEU and chrF; whitespace normalized all-page extraction"]])
    add_heading(doc, "Full Document Results")
    add_table(doc, ["Direction", "Status", "Elapsed", "BLEU", "chrF", "Structure"], rows)
    add_heading(doc, "Stability Results")
    add_table(doc, ["Direction", "Completed", "Median", "Hashes", "Stable"], stability)
    add_heading(doc, "Method and Limits")
    add_body(doc, "English-to-Cebuano output is compared with the supplied Cebuano edition; Cebuano-to-English output is compared with the supplied English original. The scores describe overlap with those editions, not a substitute for bilingual or subject-matter review. Structural checks verify a 25-page output and non-empty extractable text. Rendered pages are inspected separately for visible defects.")
    add_heading(doc, "Traceability")
    add_body(doc, "The JSON manifest, findings, and metadata inside the iteration folder contain the source/reference hashes, output hashes, service response evidence, and repeated-run samples. The supplied source and reference PDFs remain unchanged.")
    return markdown, doc


def report_comparison_and_logging(findings: dict, gpt: list[dict]) -> tuple[str, Document]:
    nllb = {row["id"]: row for row in findings["full_runs"]}
    comparison_rows = []
    for existing in gpt:
        row = nllb[existing["id"]]
        quality = row.get("quality", {})
        structure = row.get("structure", {})
        comparison_rows.append([
            existing["direction"], "NLLB", str(quality.get("bleu")), str(quality.get("chrf")), structure.get("layout_validation", "N/A").title(), minutes(row.get("elapsed_ms")),
        ])
        comparison_rows.append([
            existing["direction"], "GPT OSS", str(existing["bleu"]), str(existing["chrf"]), existing["layout_validation"].title(), "Not supplied",
        ])
    markdown = """# Part2 Iteration Comparison and Current State Logging Audit

## Outcome

This report compares the newly measured NLLB Part2 outputs with the existing GPT-OSS PDFs and audits current TriLingua logging and metric capture. GPT-OSS PDFs are scored as supplied artifacts; no GPT-OSS translation is rerun and no runtime-performance claim is made for that iteration.

## Quality and Structure Comparison

| Direction | Iteration | BLEU | chrF | Structural check | Runtime |
|---|---|---:|---:|---|---|
""" + "\n".join("| " + " | ".join(row) + " |" for row in comparison_rows) + """

The score method uses all-page extracted text after whitespace normalization. Both GPT-OSS PDFs have 25 pages and non-empty extracted text. Their file hashes are preserved in the comparison metadata. GPT-OSS model, provider settings, cache/fallback state, and elapsed time are not available from the supplied files.

## Current Logging and Metrics

- Laravel defaults to a stack containing the `single` file channel; a 14-day `daily` channel is defined but is not the default stack member.
- The document job logs failures and non-fatal persistence/notification issues, including job IDs, user IDs, filenames, exception messages, and a full trace on terminal translation failures.
- The Python service writes diagnostic messages to process output, including filename, language pair, size, provider/model, elapsed time, and exception text.
- `translation_history`, `translation_jobs`, `translation_blocks`, `translation_edit_logs`, and `translation_metrics` preserve operational and review evidence in the application database.
- `translation_metrics` stores provider/model, mode, durations, calls, block counts, cache counts, retranslation counts, and token counts when supplied.

## Audit Findings

- Filenames, user identifiers, exception messages, and stack traces can expose sensitive operational information in logs.
- Python diagnostic output can include language/source context; log calls should be audited before assuming document content never reaches a log sink.
- Laravel and Python do not currently emit a clearly shared correlation identifier that traces one request across the web app, queue worker, Python service, storage, and review records.
- The default single-file channel has no configured rotation, while the defined daily channel has a 14-day retention setting; retention is therefore not consistently enforced.
- Metrics, database audit records, and operational logs have different retention/security behavior and should not be treated as one equivalent log source.

## Recommended Next Steps

Adopt structured, correlation-aware operational events; minimize or hash file identifiers; remove full traces from routine production logs; define retention and access controls per data class; and test that source/translated text is never emitted to production logs. These are recommendations only and are not implemented by this report.
"""
    doc = new_doc("Part2 Iteration Comparison and Current State Logging Audit", "Quality and structure comparison with an audit of implemented operational evidence")
    add_body(doc, "NLLB was measured in a controlled run. Existing GPT OSS PDFs were evaluated as supplied artifacts, so the comparison is limited to quality and structure rather than speed or configuration.")
    add_heading(doc, "Quality and Structure Comparison")
    add_table(doc, ["Direction", "Iteration", "BLEU", "chrF", "Structure", "Runtime"], comparison_rows)
    add_body(doc, "Both GPT OSS artifacts contain 25 pages and extractable text. Their scores use the same all-page whitespace-normalized method as NLLB. The supplied PDFs do not include runtime configuration, cache/fallback state, model details, or elapsed time.")
    add_heading(doc, "Current Logging and Metrics")
    add_body(doc, "Laravel uses a default stack with the single file channel, while a daily 14-day channel is defined but not selected by default. Queue jobs record terminal failures and non-fatal errors. The Python translation service writes console diagnostics for document work. The database records translation history, job state, review/block evidence, edit audit data, and engine metrics.")
    add_heading(doc, "Audit Findings")
    add_table(doc, ["Area", "Observed condition", "Why it matters"], [
        ["Sensitive context", "Logs can contain filenames, user IDs, exception text, and terminal traces.", "Operational logs may expose user or document-related information."],
        ["Python output", "Service diagnostics include document/language context and errors.", "Production sinks need an explicit content-minimization review."],
        ["Correlation", "No single cross-service correlation field is evident.", "Investigating one translation across services is harder."],
        ["Retention", "Default single-file logging does not define rotation; daily logging has 14 days.", "Retention behavior is inconsistent."],
        ["Evidence types", "Logs, metrics, and review/audit tables are separate stores.", "They need distinct access and retention controls."],
    ])
    add_heading(doc, "Recommendations")
    add_body(doc, "The next logging improvement should use structured correlation-aware events, minimize or hash file identifiers, avoid routine full stack traces, define retention and access controls for each evidence type, and verify that source and translated text never reaches production logs. This audit does not change production logging.")
    return markdown, doc


def main() -> None:
    findings = load_json(ITERATION / "findings.json")
    metadata = load_json(ITERATION / "run-metadata.json")
    comparison_metadata = load_json(PART2 / "iteration-2-gptoss-artifacts.json")
    gpt = comparison_metadata["gptoss_artifacts"]

    iteration_md, iteration_doc = report_iteration(findings, metadata)
    comparison_md, comparison_doc = report_comparison_and_logging(findings, gpt)
    (PART2 / "Part2 Iteration 1 NLLB Evaluation.md").write_text(iteration_md, encoding="utf-8")
    iteration_doc.save(PART2 / "Part2 Iteration 1 NLLB Evaluation.docx")
    (PART2 / "Part2 Iteration Comparison and Current State Logging Audit.md").write_text(comparison_md, encoding="utf-8")
    comparison_doc.save(PART2 / "Part2 Iteration Comparison and Current State Logging Audit.docx")


if __name__ == "__main__":
    main()

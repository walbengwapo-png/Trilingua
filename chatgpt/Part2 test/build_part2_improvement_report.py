"""Create the evidence-based Part2 translation quality improvement report."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
PART2 = ROOT / "chatgpt" / "Part2 test"
ITERATION = PART2 / "iteration-1-nllb-200-600m"
BLUE = "17365D"
PALE = "EAF2F8"
GRID = "D9D9D9"


def load(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def minutes(value: float) -> str:
    return f"{value / 60000:.2f} min"


def set_run(run, size: float = 10.0, bold: bool = False, color: str = "000000") -> None:
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    value = OxmlElement("w:shd")
    value.set(qn("w:fill"), color)
    props.append(value)


def borders(table) -> None:
    props = table._tbl.tblPr
    value = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        child = OxmlElement(f"w:{edge}")
        child.set(qn("w:val"), "single")
        child.set(qn("w:sz"), "4")
        child.set(qn("w:color"), GRID)
        value.append(child)
    props.append(value)


def table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    item = doc.add_table(rows=1, cols=len(headers))
    item.autofit = True
    borders(item)
    for index, text in enumerate(headers):
        cell = item.rows[0].cells[index]
        shade(cell, BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(paragraph.add_run(text), 8.5, True, "FFFFFF")
    for row_index, row in enumerate(rows):
        cells = item.add_row().cells
        for index, text in enumerate(row):
            cell = cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                shade(cell, PALE)
            set_run(cell.paragraphs[0].add_run(text), 8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    paragraph.paragraph_format.space_after = Pt(4)
    set_run(paragraph.add_run(text), 13 if level == 1 else 11, True)


def body(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.1
    set_run(paragraph.add_run(text), 10.2)


def document(title: str, subtitle: str) -> Document:
    result = Document()
    section = result.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    result.styles["Normal"].font.name = "Aptos"
    result.styles["Normal"].font.size = Pt(10.2)
    title_p = result.add_paragraph(style="Title")
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(title_p.add_run(title), 20, True)
    subtitle_p = result.add_paragraph()
    subtitle_p.paragraph_format.space_after = Pt(12)
    set_run(subtitle_p.add_run(subtitle), 10.5, False, "404040")
    return result


def main() -> None:
    findings = load(ITERATION / "findings.json")
    gpt = load(PART2 / "iteration-2-gptoss-artifacts.json")["gptoss_artifacts"]
    nllb = {record["id"]: record for record in findings["full_runs"]}
    gpt_records = {record["id"]: record for record in gpt}
    evidence = []
    for identifier, direction in (("en-to-ceb", "English to Cebuano"), ("ceb-to-en", "Cebuano to English")):
        current = nllb[identifier]
        other = gpt_records[identifier]
        metrics = current["engine_response"]["metrics"]
        evidence.append([
            direction,
            str(current["quality"]["bleu"]), str(current["quality"]["chrf"]), minutes(current["elapsed_ms"]),
            str(metrics["llm_calls"]), f"{metrics['llm_total_time_ms'] / metrics['total_time_ms'] * 100:.1f}%",
            str(other["bleu"]), str(other["chrf"]),
        ])

    md = """# Part2 Translation Quality Layout and Speed Improvement Report

## Executive Finding

NLLB completed both 25-page Part2 translations with intact page counts and deterministic short-sample outputs. However, the supplied GPT-OSS artifacts have higher automatic quality scores in both directions. Engine-recorded model-call time accounts for most NLLB runtime, so the next improvement should use a quality-first mode and review workflow for user-facing documents, while optimizing the NLLB inference path separately and only retaining speed changes that pass the same quality gate.

## Evidence

| Direction | NLLB BLEU | NLLB chrF | NLLB time | NLLB calls | Model-call share | GPT-OSS BLEU | GPT-OSS chrF |
|---|---:|---:|---:|---:|---:|---:|---:|
""" + "\n".join("| " + " | ".join(row) + " |" for row in evidence) + """

Both NLLB outputs contain 25 pages and extractable text. Three repeated short samples per direction produced one output hash each. Visual checks of first, middle, and last output pages found no clipping, overlap, or page-count regression in this story-book corpus. This corpus contains no complex tables, so it does not establish table-layout quality.

## Quality Improvements

1. Keep the completed `fast` run as the speed baseline, but evaluate `balanced` as the normal user-facing NLLB path. Fast disables document analysis, document memory, semantic chunking, specialized prompts, AI quality review, cache, prepass, and layout planning. Balanced enables the first six quality features and the prepass.
2. Add a bilingual review set and score grammar, naturalness, tone, coherence across pages, terminology consistency, and meaning preservation. Use the reviewer for final acceptance, not BLEU or chrF alone.
3. Enable document memory and a glossary for recurrent names, theological terms, labels, headers, and book terminology. Save approved terminology by language pair and expose it to the review/retranslation workflow.
4. Enable AI quality review and the opt-in coherence-polish pass for Cebuano and Filipino only after it has passed the bilingual gate. Re-translate only flagged blocks, then rerun the review so a repair cannot silently worsen tone or meaning.
5. Treat titles, headers, brand names, URLs, page numbers, and source-language labels as protected or explicitly allowed passthrough spans. This prevents quality checks from either forcing awkward translations or accepting genuine untranslated body text.

## Layout Improvements

The current PDF reconstructor preserves source coordinates and tries, in order, a larger text rectangle, smaller font, and a continuation box. It can reduce text to 6 pt and a continuation box can reach the page boundary. That is a safety mechanism, not a layout-quality guarantee.

1. Add a preflight fit check before painting text: measure translated text in its original rectangle, evaluate a limited set of font and line-spacing adjustments, and choose the best fit before writing to the page.
2. Record an overflow event for every block that needs expansion, font reduction, continuation, or clipping. Fail layout QA when a continuation is clipped or a text box overlaps another content box.
3. Create a separate layout fixture set with tables, long table cells, captions, multicolumn pages, headers/footers, mixed fonts, and image-adjacent labels. The current story-book evidence cannot validate these structures.
4. Preserve the source font family and style when supported; apply a target-language fallback font only when the original cannot render the required glyphs. Use a document-specific minimum readable font rule approved by the reviewer rather than relying on the hard 6 pt fallback.
5. Add rendered-image comparison checks for page count, text-box collisions, text outside the media/crop area, and unreadable/blank text. Review all changed pages when a layout intervention occurs.

## Speed Improvements

Do not optimize extraction or reconstruction first: engine-recorded model-call time consumed 79.3% of the English-to-Cebuano runtime and 90.6% of the Cebuano-to-English runtime. NLLB currently processes one item at a time, serializes generation with a lock, and uses four beams by default.

1. Add a batched NLLB translation path that tokenizes compatible blocks together, generates in an ordered batch, and keeps per-block retry/error handling. Benchmark batch sizes 2, 4, and 8 under the same corpus and memory limit.
2. Sweep `NLLB_TORCH_THREADS` using the actual host core count, recording latency, CPU utilization, and output hashes. Keep the fastest stable setting that does not change the approved quality result.
3. Compare beam counts 4, 2, and 1 against the bilingual quality gate. A lower beam count is accepted only if it does not introduce a critical meaning/grammar/tone regression and stays within the agreed automatic-score tolerance.
4. Keep the model warm and enable the privacy-reviewed exact-block cache for production only. The controlled evaluation correctly kept cache disabled; caching improves repeated content, not first-pass translation quality.
5. Run GPT-OSS and NLLB under one identical timed protocol before choosing a production provider. GPT-OSS scored higher in the supplied artifacts, but no elapsed-time, configuration, or provider metadata was supplied, so it cannot yet be called faster or promoted on this evidence alone.

## Acceptance Gate

For each language direction and document type, release a change only when: every output page renders cleanly; no table/text box is clipped or overlaps; no protected span becomes a false translation failure; the bilingual reviewer finds no critical meaning, grammar, tone, or cohesion defect; automatic scores do not regress beyond the agreed threshold; and median latency improves or remains acceptable for the selected quality mode.

## Recommended Order

First, establish bilingual and layout fixture acceptance tests. Second, evaluate NLLB balanced mode plus glossary, memory, review, and coherence polish against the fast baseline. Third, implement and benchmark batched inference, thread tuning, and beam-count variants. Finally, run an identical GPT-OSS benchmark before any provider decision.
"""

    report = document("Part2 Translation Quality Layout and Speed Improvement Report", "Evidence-based actions for grammar tone coherence layout integrity and NLLB performance")
    body(report, "NLLB completed both Part2 translations with 25-page outputs and deterministic repeat results, but the supplied GPT OSS artifacts score higher in both directions. The recommended path is to improve translation quality and layout verification first, then optimize NLLB inference without relaxing the quality gate.")
    heading(report, "Evidence")
    table(report, ["Direction", "NLLB BLEU", "NLLB chrF", "NLLB time", "Calls", "Model calls", "GPT BLEU", "GPT chrF"], evidence)
    body(report, "Both NLLB outputs have 25 pages and non-empty text. First, middle, and final rendered pages were visually clean in this story-book corpus. The corpus has no complex tables, so it does not prove table-layout quality.")
    heading(report, "Improve Translation Quality")
    table(report, ["Priority", "Change", "Expected effect", "Validation"], [
        ["P0", "Evaluate balanced mode as the normal NLLB path", "Adds document analysis, memory, semantic chunking, prompts, review, cache, and prepass", "Compare with fast baseline using bilingual rubric"],
        ["P0", "Create bilingual grammar, tone, coherence, and terminology review set", "Makes acceptance reflect natural educational language", "No critical reviewer defect"],
        ["P1", "Use glossary and document memory for recurring terms", "Improves terminology and cross-page consistency", "Term consistency by language pair"],
        ["P1", "Use quality review and coherence polish after validation", "Targets awkward phrasing and cohesion defects", "Review after every repair"],
        ["P1", "Protect non-translatable spans", "Avoids title/header false failures and source echoes", "Span-level regression tests"],
    ])
    heading(report, "Improve Layout Quality")
    body(report, "The current PDF writer first expands a text rectangle, then reduces font size down to 6 pt, then attempts a continuation box. Add a preflight fit measurement, log every intervention, and fail layout QA on clipping or collision. Build fixtures with tables, long cells, captions, multicolumn layouts, headers/footers, and mixed fonts because this story-book corpus cannot validate them.")
    heading(report, "Improve Speed")
    body(report, "Engine-recorded model-call time consumed 79.3 percent of English-to-Cebuano runtime and 90.6 percent of Cebuano-to-English runtime. NLLB translates one block at a time, serializes generation, and uses four beams. Prioritize batched generation, a measured Torch-thread sweep, and a beam-count ablation. Keep exact-block cache for privacy-reviewed production use, not controlled first-pass evaluation.")
    heading(report, "Release Gate")
    body(report, "Release only when every page renders cleanly, no table or text box clips or overlaps, no protected span triggers a false failure, bilingual review finds no critical meaning, grammar, tone, or cohesion defect, automatic scores remain within the approved tolerance, and median latency is acceptable for the selected quality mode.")
    heading(report, "Recommended Order")
    body(report, "Establish bilingual and layout fixtures first. Then test NLLB balanced mode with glossary, memory, review, and coherence polish against fast mode. Next benchmark batching, threads, and beams. Finally run GPT OSS and NLLB with the same timed protocol before selecting a provider.")

    (PART2 / "Part2 Translation Quality Layout and Speed Improvement Report.md").write_text(md, encoding="utf-8")
    report.save(PART2 / "Part2 Translation Quality Layout and Speed Improvement Report.docx")


if __name__ == "__main__":
    main()

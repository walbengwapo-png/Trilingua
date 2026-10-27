"""Create a provider-neutral TriLingua pipeline improvement specification."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
PART2 = ROOT / "chatgpt" / "Part2 test"
ITERATION = PART2 / "iteration-1-nllb-200-600m"
BLUE = "17365D"
PALE = "EAF2F8"
GRID = "D9D9D9"


def load(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def minutes(value: float) -> str:
    return f"{value / 60000:.2f} min"


def set_run(run, size: float = 10.0, bold: bool = False, color: str = "000000") -> None:
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    value = OxmlElement("w:shd")
    value.set(qn("w:fill"), color)
    props.append(value)


def borders(table) -> None:
    props = table._tbl.tblPr
    value = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        child = OxmlElement(f"w:{edge}")
        child.set(qn("w:val"), "single")
        child.set(qn("w:sz"), "4")
        child.set(qn("w:color"), GRID)
        value.append(child)
    props.append(value)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    result = doc.add_table(rows=1, cols=len(headers))
    result.autofit = True
    borders(result)
    for index, label in enumerate(headers):
        cell = result.rows[0].cells[index]
        shade(cell, BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(paragraph.add_run(label), 8.5, True, "FFFFFF")
    for row_index, row in enumerate(rows):
        cells = result.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                shade(cell, PALE)
            set_run(cell.paragraphs[0].add_run(value), 8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    paragraph.paragraph_format.space_after = Pt(4)
    set_run(paragraph.add_run(text), 13 if level == 1 else 11, True)


def body(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.1
    set_run(paragraph.add_run(text), 10.2)


def make_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(10.2)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(title.add_run("TriLingua Provider Neutral Pipeline Improvement Report"), 20, True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    set_run(subtitle.add_run("A quality layout and speed roadmap independent of the selected translation model"), 10.5, False, "404040")
    return doc


def main() -> None:
    findings = load(ITERATION / "findings.json")
    supplied = load(PART2 / "iteration-2-gptoss-artifacts.json")["gptoss_artifacts"]
    nllb = {row["id"]: row for row in findings["full_runs"]}
    gpt = {row["id"]: row for row in supplied}
    evidence = []
    for identifier, direction in (("en-to-ceb", "English to Cebuano"), ("ceb-to-en", "Cebuano to English")):
        run = nllb[identifier]
        evidence.append([
            direction,
            f"{run['quality']['bleu']} / {run['quality']['chrf']}",
            minutes(run["elapsed_ms"]),
            f"{gpt[identifier]['bleu']} / {gpt[identifier]['chrf']}",
            "25 pages; non-empty text for both",
        ])

    markdown = """# TriLingua Provider Neutral Pipeline Improvement Report

## Decision

Improve the document-translation pipeline independently of NLLB or GPT-OSS. NLLB is not the architecture reference; it is only the reproducible local benchmark from the requested Part2 run. GPT-OSS is not a permanent default either. The pipeline must supply the same document context, terminology controls, quality gate, layout protection, scheduling, and observability to every provider. A provider is selected only after it passes the same quality, reliability, and measured-latency gate.

## Evidence

| Direction | NLLB BLEU / chrF | NLLB elapsed | Supplied GPT-OSS BLEU / chrF | Structure |
|---|---:|---:|---:|---|
""" + "\n".join("| " + " | ".join(row) + " |" for row in evidence) + """

Both NLLB outputs were 25-page, non-empty PDFs and had deterministic three-run short samples. The supplied GPT-OSS artifacts scored higher automatically in both directions, but they include no reproducible elapsed time, configuration, cache/fallback state, or provider metadata. Therefore the evidence supports GPT-OSS as a quality candidate, not an unconditional production choice.

## Target Pipeline

```text
extract PDF structure
→ build document map and reading order
→ identify protected spans and glossary terms
→ group semantic translation units with local context
→ schedule provider-neutral translation tasks
→ apply quality review and selective repair
→ preflight layout fit and reconstruct
→ render and validate
→ persist privacy-safe metrics and audit evidence
```

## Provider Neutral Interfaces

1. `TranslationUnit`: stable unit ID, page, reading order, bounding box, element role, parent section, prior/next context, table row/column headers, source text, and protected spans.
2. `ProviderCapabilities`: supported language pairs, maximum batch size, concurrency limit, context limit, streaming support, and cost/latency metadata. The scheduler reads capabilities rather than branching on provider name.
3. `translate_many(units, context, glossary)`: ordered batch result with one status/error per unit. It must preserve the unit ID even when one item fails or retries.
4. `QualityResult`: quality score, issue category, source/target snippets, repair eligibility, and reviewer provenance. It is independent from the translation provider.
5. `LayoutResult`: original/final rectangle, original/final font, expansion/reduction/continuation actions, collision status, and clipping status.

## Quality Improvements

### Context and Cohesion

- Replace isolated PDF blocks with semantic units. A paragraph keeps its heading and adjacent paragraphs; a table cell keeps its row/column header; repeated headers/footers are identified as document furniture.
- Build one document summary and terminology memory once per document. Pass only the relevant section summary, adjacent context, and glossary terms to each unit rather than repeatedly sending the entire file.
- Mark URLs, numbers, product names, page labels, brand names, and approved proper nouns as protected spans. They should survive unchanged without being mistaken for untranslated body text.
- Use a bilingual human rubric for grammar, naturalness, tone, coherence, cohesion, terminology, and meaning preservation. BLEU and chrF remain regression signals, not acceptance criteria.

### Review and Repair

- Run a provider-independent quality reviewer after initial translation. It flags missing content, source-language echoes, grammar/tone defects, inconsistent terms, and broken cross-reference context.
- Retry only the failed unit with its document context and the reviewer issue. Re-review the repaired unit and retain the original if the repair is worse.
- Offer a reviewer-approved glossary per language pair so the same educational or subject-specific term is never independently guessed on every page.

## Layout Improvements

- Preserve each source block's reading order, bounding box, alignment, font family, weight, color, and paragraph spacing where the target glyphs can render.
- Preflight every translated unit against its target rectangle before painting it. Choose an approved font/line-spacing adjustment before reconstruction instead of relying on a last-resort 6 pt minimum font or a continuation that can reach the page boundary.
- Emit a layout event whenever text expands, shrinks, continues, collides, or clips. A clipped continuation, overlap, or text outside the page box is a release failure.
- Add fixtures for tables, long table cells, columns, headings, captions, headers/footers, images, and mixed fonts. The current story-book corpus has no complex tables, so it cannot prove table quality.
- Render all changed pages and run visual regression checks for clipping, bleed/out-of-page text, collisions, blank text, and unexpected font substitution.

## Speed Improvements

- Use a capability-based scheduler: batch compatible units, obey the active provider's bounded concurrency, preserve output order, and isolate per-unit retries.
- For local sequence-to-sequence models, benchmark padded batched generation, host-specific Torch thread counts, and beam counts. Keep only variants that pass the same bilingual and layout gate.
- For remote models, limit inflight requests, combine only compatible short units, apply retry/backoff per unit, and record queue wait separately from provider latency.
- Cache only approved exact translations keyed by a privacy-safe hash of source, target, provider/model revision, glossary version, and quality profile. Do not use cache results across models or quality profiles.
- Keep model startup warm and run document analysis once per file. Do not spend effort optimizing PDF extraction/reconstruction before provider latency, because NLLB model calls dominate the measured Part2 run.

## Provider Selection Policy

Use a registry rather than a hard-coded NLLB or GPT-OSS path. Every provider must run the same corpus, quality rubric, layout fixtures, and timed protocol. Promote a provider for a quality profile only when it has no critical bilingual-review defects, passes every layout check, has reproducible configuration and metadata, meets the agreed latency target, and has an acceptable failure/retry rate.

Current recommendation: retain GPT-OSS as the leading final-quality candidate because the supplied Part2 artifacts scored higher, but do not make it mandatory or remove NLLB until GPT-OSS completes the controlled reproducibility, latency, and reliability gate. Retain NLLB as a local baseline and possible draft/offline provider, not as the pipeline design authority.

## Implementation Order

1. Introduce document-map, protected-span, and provider-capability interfaces with tests.
2. Add semantic unit grouping, section context, and glossary memory.
3. Add provider-neutral review/repair and bilingual acceptance fixtures.
4. Add layout preflight, layout events, table/column fixtures, and rendered regression checks.
5. Add capability-based batching/scheduling, cache namespace controls, and stage timing.
6. Run NLLB and GPT-OSS through the identical benchmark and select the provider per quality profile.
"""

    doc = make_document()
    body(doc, "Improve the pipeline independently of NLLB or GPT OSS. NLLB is a reproducible Part2 benchmark, not the architecture reference. GPT OSS is a quality candidate, not a permanent default. Every provider must receive the same context, terminology controls, review gate, layout protection, scheduler, and evidence requirements.")
    heading(doc, "Evidence")
    add_table(doc, ["Direction", "NLLB BLEU chrF", "NLLB time", "GPT OSS BLEU chrF", "Structure"], evidence)
    body(doc, "NLLB produced deterministic short samples and 25-page outputs. Supplied GPT OSS PDFs scored higher in both directions, but lack reproducible timing and configuration evidence. This supports a controlled provider comparison rather than an unconditional provider decision.")
    heading(doc, "Target Pipeline")
    body(doc, "Extract PDF structure, build a document map and reading order, identify protected spans and glossary terms, group semantic translation units with local context, schedule provider-neutral translation tasks, review and selectively repair, preflight layout fit, reconstruct, render, validate, and persist privacy-safe evidence.")
    heading(doc, "Provider Neutral Interfaces")
    add_table(doc, ["Interface", "Required responsibility"], [
        ["TranslationUnit", "Carries unit ID, page, reading order, box, role, section, local context, table headers, source text, and protected spans."],
        ["ProviderCapabilities", "Declares batch, concurrency, context, language, and latency capabilities for the scheduler."],
        ["translate_many", "Returns ordered per-unit status/error results without losing unit identity on retries."],
        ["QualityResult", "Records review score, issue, repair eligibility, snippets, and reviewer provenance."],
        ["LayoutResult", "Records fit actions, final geometry/font, collision state, and clipping state."],
    ])
    heading(doc, "Quality and Cohesion")
    body(doc, "Replace isolated PDF boxes with semantic units. Paragraphs retain their heading and neighboring context; table cells retain their row/column headers; recurring headers and footers are treated as document furniture. Build a document summary and terminology memory once, protect non-translatable spans, and judge grammar, tone, coherence, terminology, and meaning with a bilingual rubric. Review and selectively retry only flagged units with their context.")
    heading(doc, "Layout Integrity")
    body(doc, "Preflight translated text inside its rectangle before painting it. Preserve source styles where target glyphs support them, record every expansion/reduction/continuation, and fail on clipping, overlap, or text outside the page box. Add fixtures for tables, long cells, columns, captions, headers/footers, images, and mixed fonts; the current story-book corpus cannot validate those structures.")
    heading(doc, "Speed")
    body(doc, "Use capability-based batching and bounded concurrency, with per-unit retries and ordered output. Benchmark local batch size, Torch threads, and beam count against the same quality gate. For remote providers, measure queue wait separately, limit inflight work, and retry only affected units. Cache only privacy-safe, approved exact results in a namespace that includes provider/model, glossary version, and quality profile.")
    heading(doc, "Provider Policy")
    body(doc, "Keep GPT OSS as the leading final-quality candidate because the supplied Part2 artifacts scored higher, but do not make it mandatory until it passes a reproducible same-pipeline latency and reliability benchmark. Keep NLLB as a local baseline and possible draft/offline provider. Select each provider through the same quality, layout, reliability, and latency gate.")
    heading(doc, "Implementation Order")
    add_table(doc, ["Phase", "Deliverable"], [
        ["1", "Document map, protected spans, and provider-capability interfaces with tests."],
        ["2", "Semantic grouping, section context, and glossary memory."],
        ["3", "Provider-neutral review/repair and bilingual acceptance fixtures."],
        ["4", "Layout preflight, overflow events, table fixtures, and rendered regression checks."],
        ["5", "Capability-based scheduling, cache namespaces, and stage timing."],
        ["6", "Identical NLLB and GPT OSS benchmark and provider selection per quality profile."],
    ])

    (PART2 / "TriLingua Provider Neutral Pipeline Improvement Report.md").write_text(markdown, encoding="utf-8")
    doc.save(PART2 / "TriLingua Provider Neutral Pipeline Improvement Report.docx")


if __name__ == "__main__":
    main()

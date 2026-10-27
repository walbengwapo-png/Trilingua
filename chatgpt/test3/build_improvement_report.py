"""Document local implementation evidence; no translation provider calls."""
import json
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
OUT = HERE / "improvements"
results = json.loads((OUT / "replay-results.json").read_text(encoding="utf-8"))
verified = json.loads((OUT / "verification.json").read_text(encoding="utf-8"))
assert len(results) == 8 and all(r["validation"]["status"] == "passed" for r in results)
assert verified["originals_unchanged"]
checks = verified["tests"]["focused-tests"]
assert checks["failures"] == checks["errors"] == "0"

doc, md = Document(), []
section = doc.sections[0]
section.top_margin = section.bottom_margin = Inches(.65)
section.left_margin = section.right_margin = Inches(.7)
normal = doc.styles["Normal"]
normal.font.name, normal.font.size = "Arial", Pt(10)
normal.paragraph_format.space_after = Pt(6)
for name in ("Title", "Heading 1", "Heading 2"):
    doc.styles[name].font.name = "Arial"
    doc.styles[name].font.color.rgb = RGBColor.from_string("17365D")
section.header.paragraphs[0].text = "TriLingua | Content and layout implementation | Test3"
footer = section.footer.paragraphs[0]
footer.alignment = 2
footer.add_run("Page ")
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)


def heading(text, level=1, new_page=False):
    paragraph = doc.add_heading(text, level)
    paragraph.paragraph_format.page_break_before = new_page
    md.append("#" * (level + 1) + " " + text + "\n")


def paragraph(text):
    doc.add_paragraph(text)
    md.append(text + "\n")


def table(headers, rows):
    grid = doc.add_table(rows=1, cols=len(headers))
    grid.style = "Table Grid"
    for cell, text in zip(grid.rows[0].cells, headers):
        cell.text = text
        for run in cell.paragraphs[0].runs:
            run.bold = True
    grid.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for values in rows:
        for cell, value in zip(grid.add_row().cells, values):
            cell.text = str(value)
    for row in grid.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(4)
                for run in para.runs:
                    run.font.size = Pt(9)
    md.extend(["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"])
    md.extend("| " + " | ".join(str(v).replace("|", "/") for v in row) + " |" for row in rows)
    md.append("")


doc.add_heading("Content and Layout Improvements", 0)
md.append("# TriLingua Content and Layout Implementation Report\n")
paragraph("Date: 5 October 2026. Approved overflow policy: controlled reflow and extra pages.")
paragraph(f"Implementation completed locally. {checks['tests']} focused tests passed. All eight saved-translation PDF replays passed complete-block placement, saved-text coverage and text geometry checks. This is a rendering and guard verification result; linguistic accuracy and full end-to-end operation remain unverified.")
heading("Implemented behavior")
table(["Change", "Resulting behavior"], [
    ["Complete text and overflow", "Measure complete blocks; use a minimum 8-point translated font and bounded shrinking. Move overflow to readable continuation pages immediately after the source page. No silent word truncation in the preserved PDF route."],
    ["Rotation and geometry", "Normalize displayed coordinates for extraction and reconstruction. Transform legacy sidecars once; preserve normalized coordinates in new sidecars. Native tests cover 90, 180 and 270 degrees."],
    ["Roles, styles and artwork", "Separate heading-size changes and bullet starts. Retain mixed bold/italic runs. Avoid text, images, vector illustrations, table rules and panel edges. Protect passthrough glyphs from redaction padding. Darken pale text on white continuation pages and contrast navigation against its background."],
    ["Shared content acceptance", "Reject recognized conversational replies, empty translations and missing literal URLs/numbers/placeholders. Apply checks to primary, batch, fallback, repair, cache reuse, glossary processing and final document blocks. Keep valid initial text when a repair is invalid. Preserve language-name labels."],
    ["Saved-output gate", "Reopen the actual PDF, verify complete block placement and each saved fragment, then check bounds, orientation and text overlap. Failed checks raise a structured HTTP 422 on translation or regeneration. Successful document responses and sidecars carry validation and rendered fragment metadata."],
])
paragraph("The existing provider routes, default legacy pipeline selection and startup workflow were retained. No new dependency was installed. The PyMuPDF minimum was raised to 1.26.4 for the native PDF operations; verification used installed version 1.27.2.3. Obsolete truncation-era reconstruction helpers were removed.")

heading("Verification and document results", new_page=True)
paragraph("The replay reused each original file's existing translation and language direction. It called no translation providers and did not change originals, historical translated PDFs or the original evaluation report. Extra pages apply to outputs; the original 2–10-page input restriction remains satisfied.")
titles = ["CFPB Spending Tracker", "OSHA Working in the Heat", "CDC Inactive Tuberculosis", "FTC Avoiding Scams", "FTC Immigration Scams", "FTC Tech Support Scams", "CDC Hindi Aktibong Tuberculosis", "Ang Wikipedya"]
directions = ["EN to CEB", "EN to CEB", "EN to FIL", "FIL to CEB", "FIL to EN", "FIL to CEB", "FIL to EN", "CEB to EN"]
table(["Document", "Direction", "Pages: source / output", "Saved-PDF gate"], [
    [f"{r['id']:02d} {titles[r['id'] - 1]}", directions[r['id'] - 1],
     f"{r['validation']['source_pages']} / {r['pages']}", r['validation']['status'].upper()]
    for r in results
])
paragraph(f"Total: 23 source pages and {sum(r['pages'] for r in results)} output pages. Output PDFs total {sum(r['bytes'] for r in results) / 1_000_000:.2f} MB, compared with 39.52 MB for the historical outputs. This matched comparison isolates reconstruction; it does not measure fresh provider translation quality or latency.")
paragraph("Focused checks cover rotation, complete overflow text, numeric and URL literals, saved-PDF corruption detection, lost-block detection, redaction neighbors, vector obstacles, white panel text, bad batch responses, invalid quality/coherence repairs, glossary loss, regeneration isolation, response/sidecar metadata and structured endpoint failures. The endpoint tests use mocks; they do not exercise Laravel or live HTTP.")
paragraph("Broader run: 227 passed and 7 failed across 234 tests. The chunk-boundary failure was a Hypothesis input-generation speed check and passed in an isolated recheck. Six Gemini-analysis tests still fail against old mock response/model expectations. Their source files and the analysis/chunking implementation hashes match the saved pre-change manifest. These failures were recorded without changing the unrelated implementation or tests.")
paragraph("All current output pages were rendered and inspected in contact sheets, with full-page/crop checks of identified problems. The main CFPB icon loss was corrected. Of 24 icon regions, 20 were pixel-identical at 2x; two further regions were identical at 4x. Two regions differ at 4x, including clipped edge artwork and an adjacent rule. Exact artwork fidelity is not certified.")

heading("Limits and remaining verification", new_page=True)
paragraph("An automated PASS means the supplied target text was placed and recovered with acceptable text geometry. It does not mean the translation is correct, all source material was extracted, tables retain every semantic relationship, or the design is visually identical.")
paragraph("Six historical blocks are flagged by the new content guard: one input-request reply in file 04, four in file 05, and a missing numeric literal in file 08. These bad historical targets were deliberately retained in reconstruction stress outputs. Fresh pipeline processing rejects or retries recognized invalid content; rendering an old sidecar does not retroactively translate it. The earlier file 08 lexical mistranslation remains in its replay and is outside the literal guard's capabilities.")
paragraph("The conservative reflow produces fragmented layouts and blank original slots, especially in the spending forms and old sidecars with sentence fragments. Overflow text appears on following pages, and navigation is added only where it fits. Old source extraction and block boundaries are retained for regeneration compatibility. The saved-text gate does not certify contrast on every arbitrary background, source clipping masks, reading-order semantics, table association, PDF accessibility tags or general wrong-language detection.")
paragraph("MuPDF emits a structure-tree diagnostic for file 02 during native operations. Its output opens, renders and passes the text/geometry gate, but PDF tag/accessibility integrity remains unverified. Pre-existing clipped source artifacts are explicitly listed in validation metadata rather than treated as newly omitted translated text.")
paragraph("Remaining: a fresh live provider and Laravel upload/queue/download run; bilingual review; confirmation that the conservative continuation layout is acceptable for the forms; the six unrelated Gemini mock failures; and PDF structure/accessibility verification for file 02. No service restart, deployment, commit or production promotion was performed.")
heading("Evidence and reproduction", level=2)
paragraph("Evidence is under chatgpt/test3/improvements: replay-results.json, verification.json, focused-tests.log/XML, local-tests.log/XML, chunk-recheck.log/XML, replay.log, eight reconstructed PDFs, rendered page PNGs, contact sheets and icon-differences.png. Historical evidence remains under chatgpt/test3/evidence.")
paragraph("Offline replay: run chatgpt/test3/verify_improvements.py with the installed Python 3.11 runtime. It overwrites only the new improvements artifacts. Do not rerun run_test3.py over historical evidence. Focused test modules: test_pdf_content_layout, test_translation_fallback, test_cache_namespace, test_block_quality, test_unit_pipeline, test_regeneration and test_regression. Run with TRANSLATION_CACHE_ENABLED=false.")
paragraph("Native rotation behavior reference: https://pymupdf.readthedocs.io/en/latest/page.html#Page.remove_rotation. The source hashes, precise assertions and replay details are the authorities for the local results above.")

doc.core_properties.title = "TriLingua Test3 Content and Layout Implementation Report"
doc.save(OUT / "Test3 Content and Layout Implementation Report.docx")
(OUT / "Test3 Content and Layout Implementation Report.md").write_text("\n".join(md), encoding="utf-8")
assert len(Document(OUT / "Test3 Content and Layout Implementation Report.docx").tables) == 2
print(OUT / "Test3 Content and Layout Implementation Report.docx")

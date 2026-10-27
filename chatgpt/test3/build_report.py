"""Build the requested report from saved evidence, without provider calls."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
def read(name): return json.loads((HERE / 'evidence' / name).read_text(encoding='utf-8'))
results, manifest, qa, findings = [read(n) for n in ('results.json', 'manifest.json', 'qa.json', 'findings.json')]
assert len(results) == 8 and len({r['file'] for r in results}) == 8
doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Inches(.65)
section.left_margin = section.right_margin = Inches(.7)
style = doc.styles['Normal']
style.font.name = 'Arial'
style.font.size = Pt(10.5)
style.paragraph_format.space_after = Pt(6)
for name in ('Title', 'Heading 1', 'Heading 2'):
    doc.styles[name].font.name = 'Arial'
    doc.styles[name].font.color.rgb = RGBColor.from_string('17365D')
section.header.paragraphs[0].text = 'TriLingua | Current pipeline evaluation | Test3'
footer = section.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('Page ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
md = []
def heading(text, level=1, new_page=False):
    p = doc.add_heading(text, level)
    p.paragraph_format.page_break_before = new_page
    md.append('#' * (level+1) + ' ' + text + '\n')
def para(text):
    doc.add_paragraph(text); md.append(text + '\n')
def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Table Grid'
    for cell, value in zip(t.rows[0].cells, headers):
        cell.text = value
        for run in cell.paragraphs[0].runs: run.bold = True
    repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell, value in zip(t.add_row().cells, row): cell.text = str(value)
    for row in t.rows:
        keep = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(keep)
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(4)
                for r in p.runs:
                    r.font.name = 'Arial'
                    r.font.size = Pt(9)
    md.extend(['| ' + ' | '.join(headers) + ' |', '| ' + ' | '.join(['---']*len(headers)) + ' |'])
    md.extend('| ' + ' | '.join(str(v).replace('|', '/') for v in row) + ' |' for row in rows)
    md.append('')

doc.add_heading('Test3 Current Pipeline Evaluation', 0); md.append('# Test3 Current Pipeline Evaluation\n')
para('Run date: 5 October 2026. Eight public source PDFs; one target language per original; 23 source pages.')
para(findings['conclusion'])
para('Acceptance checks: valid output and unchanged page count; readable upright text; no visible clipping or content loss; no disruptive overlap; no conversational reply substituted for document content. FAIL below refers to these saved-output checks, not the HTTP request. Linguistic accuracy is unverified except for stated spot checks.')
table(['ID / document', 'Direction', 'Time', 'Output / QA'],
      [[f"{r['id']:02d} {findings['files'][str(r['id'])]['title']}", r['source']+' → '+r['target'], f"{r['wall_seconds']/60:.2f} min", findings['files'][str(r['id'])]['verdict']] for r in results])
para('Coverage: five of six directed language pairs. Cebuano → Filipino was not run because the collection contains only one Cebuano original. Repeating it would violate the one-translation-per-file instruction.')
heading('Scope and configuration')
para('Direct POST /translate/document on an isolated loopback FastAPI process. Balanced mode; automatic PDF column handling; current working tree; legacy block pipeline (unit flag default false). Translation: gpt-oss:20b-cloud with gemini-3.5-flash-lite fallback. Analysis/review: Gemini with Ollama fallback. Saved cache disabled for fresh processing; reload disabled; no production configuration or application source was edited.')
para('This run does not verify Laravel upload, authentication, database persistence, queue recovery, download UI, or other file formats. One startup warmup call was separate from the eight document requests. Internal batch retries and repairs were allowed as part of the current pipeline.')

heading('Observed findings', new_page=True)
for item in findings['observations']:
    heading(item['title'], 2); para(item['detail']); para('Evidence: ' + item['evidence'])

heading('Per-file content and layout checks', new_page=True)
for r in results:
    f = findings['files'][str(r['id'])]
    heading(f"{r['id']:02d}. {f['title']} ({r['source']} → {r['target']})", 2, new_page=r['id']==5)
    para(f['detail'])
    if 'structure' in r:
        s = r['structure']; m = r['metrics']
        para(f"Output: {s['page_count']} pages; {s['text_characters']:,} extracted characters; {m.get('total_blocks', 0)} blocks. Created: {r['output']}. Evidence: evidence/{r['id']:02d}_response.json and visual-qa/{r['id']:02d}_pages_*.png.")
    else: para('Request failed: ' + r.get('error', 'Unknown error'))

heading('Execution metrics and verification limits', new_page=True)
table(['ID', 'LLM calls*', 'GPT-OSS / Gemini*', 'Batch fallback', 'Echo retry', 'Repair'],
      [[r['id'], r.get('metrics', {}).get('llm_calls', '—'),
        str(r.get('metrics', {}).get('providers_used', {}).get('gptoss', {}).get('calls', 0)) + ' / ' + str(r.get('metrics', {}).get('providers_used', {}).get('gemini', {}).get('calls', 0)),
        r.get('metrics', {}).get('batch_fallbacks', '—'), r.get('metrics', {}).get('echo_retries', '—'), r.get('metrics', {}).get('retranslated_chunks', '—')] for r in results])
para('* Engine telemetry, not billing records. Provider counts describe recorded translation responses; analysis/review calls may be counted differently. See full response metrics and server.log for provenance.')
for r in results:
    if 'reference_metrics' in r:
        s = r['reference_metrics']
        para(f"File {r['id']:02d}: BLEU {s['bleu']}; chrF {s['chrf']} against the publisher's counterpart (file {'07' if r['id']==3 else '03'}). Whole-PDF text, whitespace collapsed, no content removed. These measure textual similarity; neither score is an accuracy percentage or a human evaluation.")
para(findings['limitations'])
heading('Reproducibility and next action', 2)
para('Evidence folder contains the manifest (source/engine hashes, Git revision, safe configuration), per-request response JSON, extracted source/output text, results.json, qa.json, comparison PNGs, and server.log. Originals unchanged: ' + str(manifest.get('originals_unchanged')) + '. The test process was stopped after the run; the existing Ollama service was retained.')
para('All saved engine Python source hashes were unchanged at the end of the run. Output PDFs total 39,519,216 bytes (39.52 MB), versus 3.82 MB for the originals. Source selection, download links and legal basis are documented in output/research/translation-test-resources.md.')
para(findings['next_action'])
para('Do not rerun run_test3.py over this evidence: it refuses an existing results.json to prevent duplicate translation requests. inspect_outputs.py and build_report.py only inspect/rebuild saved evidence.')
out = HERE / 'Test3 Current Pipeline Evaluation Report.docx'
doc.core_properties.title = 'TriLingua Test3 Current Pipeline Evaluation'
doc.core_properties.subject = 'Eight real document translations and saved-output inspection'
doc.save(out)
(HERE / 'Test3 Current Pipeline Evaluation Report.md').write_text('\n'.join(md), encoding='utf-8')
assert len(Document(out).tables) == 2
print(str(out))

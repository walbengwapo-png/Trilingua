from copy import deepcopy
from pathlib import Path
from hashlib import sha256
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
from docx import Document
from docx.table import Table
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json

ROOT = Path(__file__).resolve().parent.parent
REF = ROOT / 'qa/reference.docx'
OUT = ROOT / 'TriLingua_Test_Cases_Template.docx'
modules = []

def module(name, code, scenario, source):
    entry = dict(name=name, code=code, scenario=scenario, source=source, cases=[])
    modules.append(entry)
    return entry

def case(m, title, pre, steps, data, expected):
    ident = f"{m['code']}{len(m['cases'])+1:02d}"
    m['cases'].append([ident, title, pre, '\n'.join(f'{i+1}. {s}' for i,s in enumerate(steps)), data, expected, '', ''])

m = module('User Registration and Authentication', 'U-A', 'Verify registration, login, session handling and logout', 'routes/web.php; app/Http/Controllers/Auth')
case(m,'Register valid user','Logged out; email is unused.',['Open Register.','Enter data and matching confirmation.','Submit; open Dashboard.'],'Name: Test User\nEmail: user.a@example.test\nPassword: TestOnly!2026','User account is created and authenticated; Dashboard opens. Use a disposable account.')
case(m,'Reject invalid registration','Logged out.',['Submit separately with invalid name, duplicate email and weak password.'],'Name: User123\nExisting email\nPassword: abc','Each invalid field is rejected with a validation message; no account is created.')
case(m,'User login','Disposable user exists.',['Open Login.','Enter valid credentials.','Submit.'],'User A credentials','User becomes authenticated and can access Dashboard and Translate.')
case(m,'Reject invalid or empty login','Logged out; test user exists.',['Try wrong password.','Try blank fields and invalid email.'],'Wrong password; blank values; invalid email','Login is rejected; protected pages remain unavailable; validation or credential error is shown.')
case(m,'Limit failed login attempts','Isolated test IP and account.',['Submit five wrong passwords.','Submit a sixth attempt.'],'Same test email and IP','Further attempts are temporarily blocked according to the configured login limit; no protected access is granted.')
case(m,'Session refresh and logout','User A is logged in.',['Refresh Dashboard.','Log out.','Reopen a protected URL.'],'/dashboard; /history','Refresh retains the session. After logout, protected content requires login again.')
case(m,'Password reset','Reset mail is configured; disposable account exists.',['Request reset email.','Use valid link and matching new password.','Try the old password and reuse the link.'],'Test inbox; new strong password','Reset changes the password; old credentials and reused reset token are rejected. Mail delivery must be checked.')
case(m,'Google sign in','Google OAuth is configured in the test environment.',['Select Google sign in.','Complete consent.','Repeat and cancel a sign-in.'],'Authorized test Google account','Successful callback authenticates the intended user without duplicate accounts. Cancellation returns a usable login flow.')

m = module('User Dashboard and Profile', 'U-D', 'Verify user statistics, navigation and empty states', 'app/Http/Controllers/DashboardController.php; ProfileController.php; app/Services/TranslationStatsService.php')
case(m,'Dashboard empty state','User has no translations.',['Open Dashboard and Profile.','Follow Translate and History links.'],'New test user','Pages render without errors; counts reflect zero records and navigation opens the intended pages.')
case(m,'Dashboard totals','Known text and document records exist for User A.',['Count all owned records.','Open Dashboard and Profile.','Compare totals and recent activity.'],'More than 200 records; known document and word totals','Aggregates match all applicable owned records, including older records; recent activity is limited and ordered correctly.')

m = module('Text Translation Directions', 'U-T', 'Verify all six supported language directions', 'config/translation.php; app/Http/Controllers/TranslationController.php; Model/dto/requests.py')
texts = {'English':'Please submit the report tomorrow. The meeting starts at nine.', 'Filipino':'Pakisumite ang ulat bukas. Magsisimula ang pulong sa alas nuwebe.', 'Cebuano':'Palihog isumite ang taho ugma. Magsugod ang miting sa alas nuwebe.'}
for src, dst in [('English','Filipino'),('English','Cebuano'),('Filipino','English'),('Filipino','Cebuano'),('Cebuano','English'),('Cebuano','Filipino')]:
    case(m,f'{src} to {dst}','User logged in; engine and selected provider available.',['Open Translate.','Choose the language pair; enter text.','Translate; compare with a bilingual reference.'],f'{src} -> {dst}\n{texts[src]}','Output uses the target language and preserves the request, date and time meaning; history stores the pair and result. Human language review is required.')

m = module('Text Controls and Validation', 'U-V', 'Verify text actions, limits and processing modes', 'TranslationController.php; resources/views/translation.blade.php; config/translation.php')
case(m,'Processing modes','Engine and required analysis services are available.',['Translate the sample separately in each mode.','Inspect output and returned mode information.'],'fast; balanced; thorough; auto\nEnglish -> Filipino\nPlease submit the report tomorrow.','Configured modes return translated output; reported or resolved mode follows selection. Auto selects a supported mode; no exact wording equality is required.')
case(m,'Reject invalid languages and modes','Authenticated test client.',['Submit same source and target.','Submit unsupported language and mode.'],'English -> English; Spanish; mode=invalid','Invalid pair or option is rejected; no translation is stored.')
case(m,'Empty and boundary text','Record the environment text limit L. Default L is 8000 characters.',['Submit no input.','Submit L and L+1 characters separately.'],'No file or text\nL characters; L+1 characters','Empty input and L+1 are rejected. L passes length validation and is processed subject to service availability.')
case(m,'Copy download and clear','Successful text result displayed; clipboard allowed.',['Copy output and paste it.','Download text.','Clear the editor.'],'One completed text translation','Copied and downloaded text match the visible result; Clear resets the relevant text fields and errors.')
case(m,'Refresh capability contract','User logged in.',['Open Translate and its capability endpoint.','Compare options and limits.'],'/translate/capabilities','Browser and server advertise matching languages, formats, modes and limits. Defaults are 8000 characters and 51200 KB per upload; environment overrides are respected.')

m = module('Document Format Translation', 'U-F', 'Verify every accepted format and its output structure', 'config/translation.php; Model/document/extractor.py; Model/document/reconstructor.py')
formats = [
('DOCX','sample.docx with heading, bold text, table, image, header and footer','DOCX opens; translation preserves content order and inspectable document structure. Images and non-text content remain usable.'),
('PDF','sample.pdf with selectable text, heading and table','PDF opens with translated readable text; check page order, overlap and clipping against the original.'),
('TXT','sample.txt with UTF-8 text and multiple lines','TXT opens with translated text; Unicode and line structure remain readable.'),
('MD','sample.md with heading, list, link and fenced code','MD remains usable; check translated prose and intact Markdown syntax, link destinations and code.'),
('RTF','sample.rtf with paragraphs and a table','Output is DOCX, opens correctly and contains the translated content; compare formatting and content order.'),
('ODT','sample.odt with heading and table','Output is DOCX, opens correctly and contains the translated content; compare structure and formatting.'),
('CSV','sample.csv with header, quoted commas and empty cells','CSV opens; row/column counts and quoting remain valid; text is translated without shifting cells.'),
('PPTX','sample.pptx with title, body text, table and image','PPTX opens; translated text is readable in the original slide structure; compare shapes, tables and images.'),
('XLSX','sample.xlsx with text, number, date, formula and merged cells','XLSX opens; prose is translated; numbers, formulas, dates and workbook structure remain intact. Recalculate and compare formulas.'),
]
for fmt, data, expected in formats:
    case(m,f'Translate {fmt}','User logged in; valid fixture prepared; engine, worker and storage available.',['Upload fixture.','Select English -> Filipino and balanced mode.','Wait for completion; download and reopen output.'],data,expected)

m = module('PDF Columns and Document Fidelity', 'U-P', 'Verify column choices and complex document content', 'config/translation.php; Model/document/document_analyzer.py; Model/document/reconstructor.py; Model/document_translator_v3.py')
case(m,'PDF column choices','Two-column PDF has distinct left and right text.',['Translate separate copies with each column choice.','Compare text order and selected content.'],'auto; single; left; right\nDistinct column markers','Each supported choice is accepted; output follows the chosen extraction scope. Left and right select the intended column; auto detection must be checked against the fixture.')
case(m,'Scanned PDF behavior','Scanned fixture prepared; record whether OCR dependencies are installed.',['Upload image-only PDF.','Inspect completion or error.','Compare recognized text if output is produced.'],'Scanned readable English page','With working OCR, output contains translated recognized text. If OCR is unavailable or no text can be extracted, show an explicit failure; do not silently produce a blank successful file.')
case(m,'Long document consistency','Long fixture and bilingual reference prepared.',['Translate in balanced and thorough modes.','Compare first, middle and last sections.','Inspect repeated terminology and numbering.'],'Multi-page report; repeated names, dates and terms','No sections are silently omitted or duplicated; factual tokens and recurring terms remain consistent. A reviewer checks meaning, tables, page layout and overflow.')

m = module('Upload Validation and Quotas', 'U-Q', 'Verify file validation and per-user upload limits', 'config/translation.php; TranslationController.php; app/Services/QuotaService.php; DocumentsController.php')
case(m,'Reject unsupported or disguised file','Disposable user; no running upload.',['Upload unsupported file.','Rename binary file to an accepted extension and upload.'],'sample.exe; binary renamed to sample.docx','Uploads are rejected by extension or content validation; no translation job or accepted history record is created.')
case(m,'Reject corrupt document','Damaged fixture prepared.',['Upload corrupt DOCX or PDF.','Observe job and error.'],'Truncated OOXML archive or damaged PDF','Invalid content is rejected or the job fails explicitly; no successful downloadable result is advertised.')
case(m,'Upload size boundary','Record configured upload limit L; valid fixtures prepared. Default is 51200 KB.',['Upload a valid fixture at L.','Upload a valid fixture above L.'],'L KB; L+1 KB','At-limit fixture passes size validation. Above-limit fixture is rejected before translation; hosting request limits must also be recorded.')
case(m,'Daily count and byte quotas','Isolated account; record configured quotas. Defaults: 25 files and 262144000 bytes per day.',['Reach each quota independently.','Try another upload and re-translation.'],'Count limit; cumulative byte limit','Requests exceeding either quota are rejected. Re-translation consumes quota; reused or rejected duplicate work is not charged twice.')

m = module('Document Jobs and Recovery', 'U-J', 'Verify asynchronous progress, duplicates and failures', 'TranslationController.php; app/Jobs/TranslateDocumentJob.php; TranslationJob.php; resources/views/translation.blade.php')
case(m,'Queued document completion','Worker and engine are running.',['Upload a valid document.','Observe progress.','Download result; check History.'],'sample.docx; English -> Cebuano','Submission returns a job identifier; terminal completion exposes a usable result and one corresponding history record.')
case(m,'Repeated document submission','A translation is queued or processing.',['Submit the identical file and options twice rapidly.','Inspect returned identifiers and job records.'],'Same user, bytes, pair, mode and PDF setting','The active duplicate reuses the existing job; one active job and one quota reservation represent the request.')
case(m,'Refresh during processing','Document job is processing.',['Refresh Translate.','Wait for completion.','Compare job identifier.'],'Existing job identifier','The page resumes status tracking for the same job; refresh does not submit a duplicate translation.')
case(m,'Provider or engine failure','Disposable environment with controllable provider failure.',['Cause timeout, provider stop or engine unavailability.','Submit text and document jobs.','Inspect errors and terminal states.'],'Timeout; quota/authentication failure; engine offline','Failure is visible and traceable; no false successful history or download is returned. Recoverable service outages do not destroy previous translations.')
case(m,'Completion with unavailable storage','Completed record exists; storage availability can be controlled.',['Temporarily make output unavailable.','Poll status and request download.','Restore storage; retry.'],'Missing file and temporary storage outage tested separately','Download unavailability is reported honestly; no broken ready link is shown. Temporary outage and missing data are distinguished; restored output becomes downloadable.')

m = module('Translation History and Bookmarks', 'U-H', 'Verify owned history, actions and record navigation', 'HistoryController.php; BookmarksController.php; app/Services/HistoryService.php')
case(m,'History and pagination','User has text and document history, including more than 50 records.',['Open History.','Navigate pages.','Open text and document details.'],'Known record IDs and timestamps','Owned records are accessible through pagination; pair, type, status and result match each record; details open the selected record.')
case(m,'Rename bookmark and priority','User owns a disposable record.',['Rename it.','Add and remove bookmark.','Toggle priority review.','Refresh related lists.'],'Name: Reviewed report','Changes persist for the selected owned record; Bookmarks reflects the toggle and priority state is visible to review workflows.')
case(m,'Original and translated downloads','Completed owned document has original and output files.',['Download original and translated file.','Open previews and files.'],'Known source and result hashes','Correct files are returned with usable names and types; original content remains unchanged. Previews render safely.')
case(m,'Delete history record','Disposable owned document; identify any linked translations.',['Delete selected record.','Refresh History, Documents and Bookmarks.','Try its former download links.'],'Disposable record and related file references','Deleted record disappears from owned lists and cannot be accessed. Cleanup must preserve files still referenced by other records or active jobs.')

m = module('My Documents and Re Translation', 'U-M', 'Verify grouping, filters and stored-original translation', 'DocumentsController.php; resources/views/my-documents.blade.php; app/Services/HistoryService.php')
case(m,'Document grouping and filters','Multiple originals and translations span more than one page.',['Open My Documents.','Use name, pair, date and type filters.','Navigate pages and clear filters.'],'Different names, language pairs, dates and types','Page shows owned originals with the correct translations; filters behave consistently within their displayed scope and older records remain reachable through pagination.')
case(m,'Translate original to another language','Owned English original and Filipino result exist in storage.',['Request Cebuano from the existing original.','Wait for completion.','Open its grouped result.'],'Original English document; target Cebuano','Stored original is reused; completed child is linked to the correct source document and is downloadable.')
case(m,'Repeated re-translation','Owned original exists; same target is processing or completed.',['Repeat the same re-translation request.','Inspect job and returned result.'],'Same original and target','Active work is reused; a completed matching result is reused when available. No duplicate provider work or quota charge is created.')
case(m,'Unavailable original','Control original-file storage in test environment.',['Request re-translation with missing original.','Repeat during temporary storage outage.'],'Missing original; unknown storage availability','Missing original prompts re-upload. Temporary availability uncertainty prompts retry; it is not reported as permanent loss.')

m = module('Notifications and Account Settings', 'U-S', 'Verify notification state, account changes and themes', 'NotificationController.php; SettingsController.php; ProfileController.php')
case(m,'Notifications','User has read and unread test notifications.',['Open Notifications and unread filter.','Follow an owned translation link.','Mark one then all as read.'],'Known owned notification IDs','Notifications link to the intended records; unread state and count update correctly and persist after refresh.')
case(m,'Update account details','Disposable User A; User B has a different email.',['Update name and unused email.','Refresh Profile.','Attempt duplicate email.'],'Test User Updated; unused email; User B email','Valid changes persist; duplicate or invalid email is rejected; another account is not modified.')
case(m,'Change password','Disposable user knows current password.',['Submit wrong current password.','Submit valid current and matching new password.','Log in with old then new password.'],'Test-only current and new passwords','Wrong current password is rejected. Valid change succeeds; old password fails and new password authenticates.')
case(m,'Light and dark themes','User logged in.',['Save light then dark theme.','Open Translate, History and Settings.','Refresh.'],'theme=light; theme=dark','Chosen theme persists and remains readable across pages; controls and messages maintain usable contrast.')

m = module('Admin Access and Analytics', 'A-D', 'Verify admin access, aggregates and CSV exports', 'routes/web.php; app/Http/Controllers/Admin/DashboardController.php; app/Http/Middleware/EnsureUserIsAdmin.php')
case(m,'Admin login and access','Provisioned test admin exists.',['Log in with admin credentials.','Open /admin and each admin navigation item.'],'Test admin credentials','Admin pages open for the authorized role; navigation and data render without errors.')
case(m,'Admin dashboard aggregates','Known pending, verified, edited and flagged records exist.',['Open Admin Dashboard.','Apply supported date controls.','Compare totals and trends with fixture records.'],'Known review counts, dates and metrics','Dashboard summaries, trends and health information reflect their applicable records and filters; absent metrics are not presented as measured performance.')
case(m,'CSV exports','Admin and known review/user/flag fixtures exist.',['Export review trends, flags and users.','Reopen CSV files.','Compare rows with source records.'],'Three export endpoints; cells beginning with =, +, -, @','Exports open with correct headers and applicable data. Spreadsheet formula-like user input is escaped safely.')

m = module('Admin Review Queue and Text Review', 'A-T', 'Verify review selection, text correction, flags and verification', 'Admin/ReviewController.php; Admin/TextReviewController.php; app/Services/Admin/ReviewService.php')
case(m,'Review queue and filters','Admin; mixed translation types and review states.',['Open Review queue.','Apply visible search/status/type controls.','Open selected item.'],'Text/document; pending/verified/edited/flagged; priority item','Filters show matching records; selected review opens the correct original, output and review metadata.')
case(m,'Verify text translation','Admin; pending text item.',['Compare original and translation.','Verify.','Refresh review and owner History.'],'Known pending text ID','Review status becomes verified and an audit entry identifies the actor and action; owner sees the updated status.')
case(m,'Edit text translation','Admin; text item with known output.',['Change translated text and add note.','Save.','Compare owner result and audit trail.'],'Corrected target text; review note','Corrected text persists, review status reflects editing, and the audit trail retains the change; owner receives the saved result.')
case(m,'Flag text translation','Admin; text item.',['Enter supported reason and note.','Flag.','Inspect queue and audit trail.'],'Meaning error; reviewer note','Flag reason, status and note persist and are auditable; invalid or missing required reason is rejected.')

m = module('Admin Document Draft and Publication', 'A-P', 'Verify safe edits, regeneration, verification and concurrent review', 'Admin/DocumentReviewController.php; app/Services/Admin/ReviewService.php; Model/document/regenerator.py')
case(m,'Open document review','Admin; document with original, output, sidecar and blocks.',['Open review.','Inspect original, translated file and ordered blocks.'],'Known DOCX document and block IDs','Review displays the correct original and published translation; block order and text match the document.')
case(m,'Save block as draft','Admin; editable document block; baseline output saved for comparison.',['Edit a block and save draft.','Reload editor.','Download as owner.'],'Corrected text; current draft revision','Draft persists and revision advances; owner download remains the last published version until regeneration publishes the edit.')
case(m,'Publish regenerated file','Admin; current draft and storage pointer.',['Use Save and Regenerate.','Reopen returned file and owner download.','Compare draft and published revisions.'],'Current block draft; valid regeneration dependencies','A readable new file contains the edits; publication updates file pointer and published block text atomically; original remains available.')
case(m,'Verify and flag document','Admin; first a document with unpublished edits, then a published document.',['Try Verify before publication.','Publish edits and Verify.','Flag a separate document with reason.'],'Unpublished draft; published document; flag reason','Unpublished edits prevent verification. Published content can be verified; flag status/reason persist with audit entries.')
case(m,'Concurrent admin edits','Two test admins open the same document revision.',['Admin A saves an edit.','Admin B saves stale revision.','Repeat with competing publication.'],'Same block; stale draft revision or storage pointer','Stale writes return a conflict and prompt reload; newer draft and published content are not overwritten.')
case(m,'Failed regeneration or storage','Admin; published file exists; controlled engine/storage failure.',['Save a draft.','Cause regeneration or readback failure.','Attempt publication; download prior output.'],'Engine error; failed upload/readback','Publication fails explicitly; prior published file and pointer remain usable; draft is retained and orphan cleanup is tracked.')

m = module('Admin Jobs User Directory and Audit', 'A-O', 'Verify operational monitoring and read-only directories', 'Admin/JobController.php; Admin/UserController.php; Admin/AuditLogController.php')
case(m,'Job monitoring and retry','Admin; recoverable failed document job with original payload.',['Open Jobs; filter by status/search.','Retry recoverable failure once.','Inspect completion and repeated retry.'],'Known failed job; retained source file','Retry requeues eligible work without duplicate completion; missing payload or ineligible retry is refused clearly.')
case(m,'User directory and totals','Admin; users with known translations, including more than 200.',['Search and sort users.','Open user overview and translations.','Apply available filters and pagination.'],'Name/email; recent/name/translations sort','Directory remains read-only; correct user records and full aggregates appear; filters and pagination reach applicable translations.')
case(m,'Audit trail','Login, account change, review edit, flag and publish fixtures exist.',['Open Audit.','Filter review/account type, action and date.','Inspect actor, target and change details.'],'Known audit IDs and dates','Entries match recorded events and actors; filters return relevant actions without exposing passwords or service secrets.')

m = module('Security Usability and Non Functional Checks', 'N-X', 'Verify isolation, safe previews, accessibility and measured behavior', 'routes/web.php; HistoryController.php; NotificationController.php; Model/server.py; resources/views')
case(m,'Guest and role restrictions','Logged-out session and regular User A.',['As guest, open protected URLs.','As User A, request admin pages and write actions.'],'/dashboard; /history; /admin; admin review POST','Guests must authenticate; regular users cannot read admin data or perform admin mutations.')
case(m,'Ownership isolation','Users A and B own different files, jobs and notifications.',['As A, request B detail, file, status and actions.','Attempt rename, bookmark, delete and re-translation for B.'],'User B record/job/notification IDs','Cross-user access is denied or treated as unavailable; B records, files and notification state are unchanged.')
case(m,'Safe document previews','Fixtures contain HTML-like text and script payloads.',['Preview original and translated documents.','Open history details and review blocks.'],'Text: <script>alert(1)</script>\nHTML event handlers','Untrusted document text renders safely; scripts or event handlers do not execute in the application origin.')
case(m,'CSRF and engine authorization','Authenticated client; configured engine service token.',['Submit write request without CSRF token.','Call protected engine endpoints without and with invalid service token.'],'Missing CSRF; missing/wrong X-Service-Token','Unauthorized writes and engine calls are rejected; no translation, review or account mutation occurs.')
case(m,'Responsive keyboard and empty states','Desktop and mobile browsers; keyboard-only session.',['Open main user and admin pages at 360 and 1280 px widths.','Navigate forms/dialogs by keyboard.','Inspect empty and error states.'],'360 px; 1280 px; Tab/Enter/Escape','Required actions remain reachable; focus and labels are usable; no clipped controls; empty/error states explain the next action.')
case(m,'Performance and concurrent work','Record hardware, provider, fixture size and agreed acceptance limits before running.',['Measure page loads and text/document completion.','Submit concurrent work from two users.','Inspect states, totals and outputs.'],'Small and long documents; two users; approved timing limits','Record measured timings against the approved limits. Work stays associated with its owner; no duplicate or lost completion. No timing threshold is assumed by this template.')
case(m,'Human translation acceptance','Bilingual reviewers and agreed reference corpus/rubric prepared.',['Review six directions and all relevant formats.','Check meaning, terms, names, numbers and omissions.','Record errors against the approved rubric.'],'Prepared bilingual references; approved quality criteria','Translations meet the agreed human quality criteria; automated scores alone do not establish acceptance.')

ref_hash = sha256(REF.read_bytes()).hexdigest()
d = Document(REF)
metadata_xml = deepcopy(d.tables[1]._tbl)
cases_xml = deepcopy(d.tables[2]._tbl)
concerns_xml = deepcopy(d.tables[0]._tbl)
section_xml = deepcopy(d.sections[0]._sectPr)
body = d._element.body
for child in list(body):
    body.remove(child)
body.append(section_xml)

def clean(element):
    # Google export contains fractional twip values and revision remnants.
    # Normalize measurements only in cloned components used in this template.
    for e in list(element.iter()):
        if etree.QName(e).localname.endswith('Change'):
            e.getparent().remove(e)
    for e in element.iter():
        for key in list(e.attrib):
            if etree.QName(key).localname in ('w','val') and '.' in e.attrib[key]:
                try:
                    e.attrib[key] = str(round(float(e.attrib[key])))
                except ValueError:
                    pass

def fill_cell(cell, text):
    p = cell.paragraphs[0]
    rpr = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    for child in list(cell._tc):
        if child.tag != qn('w:tcPr'):
            cell._tc.remove(child)
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.keep_together = False
    run = p.add_run(text)
    if rpr is not None:
        run._r.insert(0,rpr)
    # Retain template Arial 12 pt without copied links or emphasis.
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    run.bold = False
    run.italic = False
    run.font.color.rgb = RGBColor(0,0,0)

def insert(element):
    body.insert(len(body)-1,element)

title = d.add_paragraph('TriLingua Test Cases Template',style='Title')
title.paragraph_format.space_after = Pt(14)
for r in title.runs:
    r.font.color.rgb = RGBColor(0,0,0)
    r.font.name = 'Arial'
    r.font.size = Pt(22)
d.add_paragraph('Test inputs and expected outcomes are prepared for TriLingua. Complete the environment and tester fields before execution. Leave Actual Result and Status blank until a test is run; enter observed evidence and Pass, Fail or Blocked afterward.')
d.add_paragraph('Scope and preparation',style='Heading 1')
d.add_paragraph('User and Admin workflows cover English, Filipino and Cebuano in all six directions. Accepted inputs are DOCX, PDF, TXT, MD, RTF, ODT, CSV, PPTX and XLSX. RTF and ODT produce DOCX; other configured formats retain their extension. Processing choices are fast, balanced, thorough and auto. PDF column choices are auto, single, left and right.')
d.add_paragraph('Default validation limits are 8000 text characters and 51200 KB per upload. Daily document quotas default to 25 files and 262144000 bytes per user. Record actual environment overrides before boundary tests. Provider, OCR, mail, Google OAuth, storage and queue tests require their corresponding services to be configured.')
d.add_paragraph('Use disposable users and documents for mutation, deletion, quota and failure tests. Sample addresses and passwords are test inputs, not existing accounts. Prepare each named fixture before execution; these fixtures are not included in this template. Record the build revision, provider, browser and service configuration in Environment. Agree performance and language-quality criteria before acceptance tests.')
d.add_paragraph('Coverage',style='Heading 1')
for i in range(0,len(modules),3):
    d.add_paragraph('   |   '.join(f"{m['name']} ({len(m['cases'])})" for m in modules[i:i+3]))
d.add_paragraph(f"Prepared 12 October 2026   |   {sum(len(m['cases']) for m in modules)} test cases   |   No execution results recorded")
for p in d.paragraphs:
    if p.style.name == 'Title':continue
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.05
    for run in p.runs:
        run.font.name = 'Arial'
        if p.style.name not in ('Heading 1','Heading 2'):
            run.font.size = Pt(12)

for m in modules:
    d.add_page_break()
    h = d.add_paragraph(m['name'],style='Heading 1')
    h.paragraph_format.space_after = Pt(8)
    meta_el = deepcopy(metadata_xml)
    clean(meta_el)
    insert(meta_el)
    meta = Table(meta_el,d._body)
    left = meta.cell(0,0).tables[0]
    for i,val in enumerate(['TriLingua',m['name'],'','','12 October 2026','']):
        fill_cell(left.cell(i,1),val)
    right = meta.cell(0,1).tables[0]
    fill_cell(right.cell(0,1),m['cases'][0][0]+' to '+m['cases'][-1][0])
    fill_cell(right.cell(1,1),m['scenario'])
    # Preserve labels, blue fill, borders, nested metadata layout and widths.
    for tbl in [meta,left,right]:
        for row in tbl.rows:
            trpr = row._tr.get_or_add_trPr()
            for height in list(trpr.findall(qn('w:trHeight'))):
                trpr.remove(height)
    spacer=d.add_paragraph()
    spacer.paragraph_format.space_after = Pt(8)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = 1
    table_el = deepcopy(cases_xml)
    clean(table_el)
    trs=table_el.findall(qn('w:tr'))
    row_pattern=deepcopy(trs[1])
    for tr in trs[1:]: table_el.remove(tr)
    header=trs[0]
    repeat=OxmlElement('w:tblHeader')
    header.get_or_add_trPr().append(repeat)
    for row_data in m['cases']:
        row_el=deepcopy(row_pattern)
        trpr=row_el.get_or_add_trPr()
        for height in list(trpr.findall(qn('w:trHeight'))):trpr.remove(height)
        no_split=OxmlElement('w:cantSplit')
        trpr.append(no_split)
        table_el.append(row_el)
    insert(table_el)
    table=Table(table_el,d._body)
    for row, row_data in zip(table.rows[1:],m['cases']):
        for cell_,value in zip(row.cells,row_data):fill_cell(cell_,value)
    for c in table.rows[0].cells:
        for p in c.paragraphs:
            p.paragraph_format.keep_with_next=True
            p.paragraph_format.space_after=Pt(0)

d.add_page_break()
d.add_paragraph('Identified Concerns',style='Heading 1')
d.add_paragraph('Record findings only after executing tests. Link each finding to the relevant test case and attach evidence separately.')
clean(concerns_xml)
trs=concerns_xml.findall(qn('w:tr'))
pattern=deepcopy(trs[1])
for tr in trs[1:]:concerns_xml.remove(tr)
for _ in range(5):concerns_xml.append(deepcopy(pattern))
insert(concerns_xml)
t=Table(concerns_xml,d._body)
for row in t.rows[1:]:
    for cell_ in row.cells:fill_cell(cell_,'')
    row.height=Pt(35)
d.add_page_break()
d.add_paragraph('Capability sources',style='Heading 1')
d.add_paragraph('TriLingua source review date 12 October 2026. Capability inputs were checked in the current working copy, including local changes. Test expectations are acceptance criteria and are not evidence of successful execution.')
d.add_paragraph('Reference format: PrimeraAI TestCases, Google document 1UTsKchbl8YzRzGTKWecYd1rxS1eQrXM3V0p6VxrlbtI. The role/module metadata, eight test columns, blue headers and legal landscape page format are retained for TriLingua.')
for m in modules:
    p=d.add_paragraph(f"{m['code']}  {m['source']}")
    for r in p.runs:r.font.size=Pt(9)

# Patch only body XML; other reference package parts remain byte-identical.
document_xml=etree.tostring(d._element,xml_declaration=True,encoding='UTF-8',standalone=True)
with ZipFile(REF) as zin, ZipFile(OUT,'w',ZIP_DEFLATED) as zout:
    for info in zin.infolist():
        data=document_xml if info.filename=='word/document.xml' else zin.read(info.filename)
        zout.writestr(info,data)

with ZipFile(REF) as zin,ZipFile(OUT) as zout:
    assert all(zin.read(n)==zout.read(n) for n in zin.namelist() if n!='word/document.xml')
check=Document(OUT)
test_tables=[t for t in check.tables if len(t.columns)==8]
actual_rows=[r for t in test_tables for r in t.rows[1:]]
ids=[r.cells[0].text for r in actual_rows]
assert len(ids)==len(set(ids))==sum(len(m['cases']) for m in modules)
assert all(r.cells[6].text==r.cells[7].text=='' for r in actual_rows)
assert all(all(r.cells[i].text.strip() for i in range(6)) for r in actual_rows)
assert 'PrimeraAI' not in ''.join(c.text for t in check.tables for r in t.rows for c in r.cells)
assert sha256(REF.read_bytes()).hexdigest()==ref_hash
assert len(check.sections)==1

contract=f'''# Template contract
Reference: {REF}
SHA256: {ref_hash}
Reference PDF: qa/reference-google.pdf, 51 pages; selected role/module table pattern inspected on page 6.
Page system: legal landscape 14 by 8.5 inches, margins 0.5 inches, one final section.
Typography: Arial 12 pt table content; template white bold headers on blue fill; 22 pt black document title.
Tables: retain nested two-column metadata, six left metadata rows and two right rows; original eight test columns and widths.
Slots: project, module, test identifiers, scenario, preconditions, procedures, sample data and expected results rewritten for TriLingua. Environment, Created By and Date of Review blank. Actual Result and Status empty in every case. Creation date 12 October 2026.
Intentional adaptations: use User and Admin roles; remove unrelated PrimeraAI findings and results; clone pattern across current modules; normalize fractional twips and remove revision remnants only in cloned elements; auto row heights, repeated headers, no split data rows; add preparation and source pages. Source PDF dropdowns become blank cells in Word output.
Preservation: every OOXML package part except word/document.xml remains byte-identical. Original reference remains unchanged. New content uses cloned actual source tables rather than a new design.
Final checks: unique IDs, complete planned columns, blank actual/status cells, no foreign project data in tables, preserved page geometry and package parts; visual review of all final pages required.
'''
(ROOT/'qa/artifact.md').write_text(contract,encoding='utf-8')
(ROOT/'qa/cases.json').write_text(json.dumps(modules,indent=2,ensure_ascii=False),encoding='utf-8')
print(json.dumps({'output':str(OUT),'modules':len(modules),'cases':len(ids),'blank_actual_and_status':True,'reference_unchanged':True}))

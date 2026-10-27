from pathlib import Path
import re
from urllib.parse import quote
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).parent
SOURCE = Path('C:/dev/Trilingua/docs/CODEBASE_WALKTHROUGH_2026-10-15.md')
OUTPUT = Path('C:/Users/walbe/Downloads/TriLingua_Codebase_Walkthrough.docx')

def diagram():
    im = Image.new('RGB', (1800, 850), 'white')
    d = ImageDraw.Draw(im)
    font = ImageFont.truetype('C:/Windows/Fonts/segoeui.ttf', 32)
    small = ImageFont.truetype('C:/Windows/Fonts/segoeui.ttf', 26)
    boxes = {
        'browser': (30, 35, 510, 190),
        'laravel': (660, 35, 1140, 190),
        'work': (1290, 35, 1770, 190),
        'store': (30, 340, 510, 495),
        'python': (660, 340, 1140, 495),
        'bridge': (1290, 340, 1770, 495),
        'translate': (350, 650, 830, 805),
        'review': (980, 650, 1460, 805),
    }
    labels = {
        'browser': 'Browser\nText or uploaded file',
        'laravel': 'Laravel controller\nValidation and acceptance',
        'work': 'Text request or\nqueued document worker',
        'store': 'Laravel persistence\nFiles history and blocks',
        'python': 'Python service\nTranslation pipelines',
        'bridge': 'TranslationManager\nHTTP bridge',
        'translate': 'Translation provider\nGPT OSS and fallback',
        'review': 'Analysis and review\nSeparate provider chain',
    }
    def arrow(start, end, label=None):
        d.line([start, end], fill='#555555', width=5)
        import math
        a = math.atan2(end[1]-start[1], end[0]-start[0])
        p1=(end[0]-21*math.cos(a-0.4),end[1]-21*math.sin(a-0.4))
        p2=(end[0]-21*math.cos(a+0.4),end[1]-21*math.sin(a+0.4))
        d.polygon([end,p1,p2],fill='#555555')
        if label:
            d.text(((start[0]+end[0])/2,(start[1]+end[1])/2-40),label,font=small,fill='#333333',anchor='mm')
    arrow((510,112),(660,112))
    arrow((1140,112),(1290,112))
    arrow((1530,190),(1530,340))
    arrow((1290,417),(1140,417))
    arrow((660,417),(510,417))
    arrow((270,340),(270,190),'Result')
    arrow((790,495),(590,650))
    arrow((1010,495),(1220,650))
    for k,box in boxes.items():
        d.rounded_rectangle(box,radius=12,fill='#F4F6F8',outline='#76818C',width=2)
        x=(box[0]+box[2])/2; y=(box[1]+box[3])/2
        d.multiline_text((x,y),labels[k],font=font,fill='black',anchor='mm',align='center',spacing=10)
    out=ROOT/'architecture.png'; im.save(out)
    return out

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.65)
sec.left_margin=Inches(.7); sec.right_margin=Inches(.7)
sec.footer_distance=Inches(.28)
for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet','List Number'):
    s=doc.styles[name]; s.font.name='Calibri'; s.font.color.rgb=RGBColor(0,0,0)
for s in doc.styles:
    for border in s.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
normal=doc.styles['Normal']; normal.font.size=Pt(10.5)
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.line_spacing=1.08
for name,size,before,after in [('Title',25,0,13),('Heading 1',15,15,7),('Heading 2',12,10,5),('Heading 3',11,8,4)]:
    s=doc.styles[name]; s.font.size=Pt(size); s.font.bold=True
    s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after)
    s.paragraph_format.keep_with_next=True
    s.paragraph_format.widow_control=True
doc.styles['Subtitle'].font.size=Pt(11)
code_style=doc.styles.add_style('Code Block',1)
code_style.font.name='Consolas'; code_style.font.size=Pt(9)
code_style.paragraph_format.space_after=Pt(2)
code_style.paragraph_format.line_spacing=1.03
code_style.paragraph_format.left_indent=Inches(.12)

token=re.compile(r'(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`)')
def inline(p,text):
    for part in token.split(text):
        if not part: continue
        m=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',part)
        if m:
            label,target=m.groups()
            lm=re.search(r':(\d+)$',target)
            if lm and target.startswith('C:/'):
                label=re.sub(r':'+lm.group(1)+r'$','',label)
                label += ' line '+lm.group(1); target=target[:lm.start()]
            if target.startswith('C:/'): target='file:///'+quote(target,safe='/:')
            rel=p.part.relate_to(target,RT.HYPERLINK,is_external=True)
            h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),rel)
            r=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
            color=OxmlElement('w:color'); color.set(qn('w:val'),'245580'); pr.append(color)
            r.append(pr); t=OxmlElement('w:t'); t.text=label; r.append(t); h.append(r); p._p.append(h)
        else:
            r=p.add_run(part[2:-2] if part.startswith('**') else part[1:-1] if part.startswith('`') else part)
            if part.startswith('**'): r.bold=True
            elif part.startswith('`'): r.font.name='Consolas'; r.font.size=Pt(9.5)

def make_table(rows):
    cols=len(rows[0]); table=doc.add_table(rows=0,cols=cols)
    table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
    widths=([1.3,2.5,3.17] if cols==3 else [2.25,4.72])
    # Restrict the sum to the 6.87 inch text area.
    factor=6.87/sum(widths); widths=[x*factor for x in widths]
    for col,w in zip(table.columns,widths): col.width=Inches(w)
    pr=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    pr.append(borders)
    for ri,row in enumerate(rows):
        cells=table.add_row().cells
        trpr=table.rows[-1]._tr.get_or_add_trPr()
        split=OxmlElement('w:cantSplit'); trpr.append(split)
        if ri==0:
            repeat=OxmlElement('w:tblHeader'); trpr.append(repeat)
        for ci,(cell,value) in enumerate(zip(cells,row)):
            cell.width=Inches(widths[ci]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
            tcpr.append(margins)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E3EAF0');tcpr.append(shade)
            p=cell.paragraphs[0]; inline(p,value)
            p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1)
            p.paragraph_format.line_spacing=1.03
            p.paragraph_format.keep_with_next=False
            for r in p.runs: r.font.size=Pt(9.3);r.bold=ri==0 or r.bold
            for rp in p._p.xpath('.//w:rPr'):
                for sz in rp.findall(qn('w:sz')): rp.remove(sz)
                sz=OxmlElement('w:sz');sz.set(qn('w:val'),'19');rp.append(sz)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

lines=SOURCE.read_text(encoding='utf-8').splitlines()
i=0; diagram_path=diagram()
while i<len(lines):
    line=lines[i]
    if not line.strip(): i+=1;continue
    if line.startswith('Result: current-source walkthrough'): break
    if line.startswith('```'):
        language=line[3:].strip(); content=[];i+=1
        while i<len(lines) and not lines[i].startswith('```'):
            content.append(lines[i]);i+=1
        if language=='mermaid':
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(diagram_path),width=Inches(6.8))
            pic=p._p.xpath('.//wp:docPr')[0];pic.set('descr','Browser requests pass through Laravel and TranslationManager to Python. Python uses separate translation and analysis providers. Results are stored by Laravel and returned to the browser.')
        else:
            for ci,cl in enumerate(content):
                p=doc.add_paragraph(style='Code Block');p.add_run(cl)
                p.paragraph_format.keep_with_next=ci<len(content)-1
        i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            values=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?',x) for x in values): rows.append(values)
            i+=1
        make_table(rows);continue
    m=re.match(r'^(#{1,3})\s+(.+)',line)
    if m:
        level=len(m.group(1)); heading=m.group(2)
        heading=re.sub(r'[^\w\s]',' ',heading);heading=re.sub(r'\s+',' ',heading).strip()
        p=doc.add_paragraph(style='Title' if level==1 else 'Heading '+str(level-1));inline(p,heading)
        i+=1;continue
    if line.startswith('- '):
        p=doc.add_paragraph(style='List Bullet');inline(p,line[2:]);i+=1;continue
    if re.match(r'^\d+\. ',line):
        p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.2);p.paragraph_format.first_line_indent=Inches(-.2)
        inline(p,line);i+=1;continue
    parts=[line];i+=1
    while i<len(lines) and lines[i].strip() and not re.match(r'^(#|\||```|- |\d+\. )',lines[i]):
        parts.append(lines[i]);i+=1
    p=doc.add_paragraph();inline(p,' '.join(parts))

footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('TriLingua walkthrough  |  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor(85,85,85)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='TriLingua codebase walkthrough'
doc.core_properties.subject='Data flow translation logic model selection and feature file map'
doc.core_properties.author=''
doc.core_properties.keywords='TriLingua Laravel Python translation codebase'
doc.save(OUTPUT)
print(OUTPUT)
print('Paragraphs',len(doc.paragraphs),'Tables',len(doc.tables))

"""Formula-preservation regression checks; no model or storage requests."""
import sys
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.formula import ArrayFormula

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from document.extractor import read_xlsx
from document.reconstructor import translate_xlsx_inplace
from document.regenerator import reconstruct_from_sidecar


@pytest.fixture
def workbook(tmp_path):
    path = tmp_path / 'source.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.title = 'Lesson'
    ws['A1'] = '=SUM(B1:B2)'
    ws['B1'] = 2
    ws['B2'] = 3
    ws['C1'] = 'Read the worksheet'
    ws['C1'].font = Font(bold=True)
    ws['D1'] = '=literal text'
    ws['D1'].data_type = 's'
    ws['E1'] = ArrayFormula(ref='E1:E2', text='=B1:B2*2')
    ws.merge_cells('G1:H1')
    ws['G1'] = 'Answer here'
    wb.create_sheet('Answers')['A1'] = '=Lesson!A1'
    wb.save(path)
    return path


def assert_formulas(path):
    wb = load_workbook(path, data_only=False)
    assert wb['Lesson']['A1'].value == '=SUM(B1:B2)'
    assert wb['Answers']['A1'].value == '=Lesson!A1'
    assert wb['Lesson']['E1'].value.text == '=B1:B2*2'
    assert wb['Lesson']['E1'].value.ref == 'E1:E2'
    assert wb['Lesson']['B1'].value == 2
    assert wb['Lesson']['B2'].value == 3
    assert 'G1:H1' in wb['Lesson'].merged_cells
    assert wb['Lesson']['C1'].font.bold
    wb.close()


def test_initial_translation_retains_formulas_and_literal_formula_like_text(workbook, tmp_path):
    output = tmp_path / 'translated.xlsx'
    seen = []
    def translate(text, **kwargs):
        seen.append(text)
        return '=translated literal' if text.startswith('=') else 'Translated: ' + text
    translate_xlsx_inplace(workbook, output, translate)
    assert_formulas(output)
    wb = load_workbook(output)
    assert wb['Lesson']['D1'].data_type == 's'
    assert wb['Lesson']['D1'].value == '=translated literal'
    assert seen == ['Read the worksheet', '=literal text', 'Answer here']
    wb.close()


def test_extraction_excludes_standard_cross_sheet_and_array_formulas(workbook):
    assert [b['text'] for b in read_xlsx(workbook)] == [
        'Read the worksheet', '=literal text', 'Answer here']


def test_regeneration_retains_formulas_and_applies_edited_text(workbook, tmp_path):
    output = tmp_path / 'regenerated.xlsx'
    sidecar = {'format': '.xlsx', 'blocks': [
        {'block_index': 0, 'text': 'Edited worksheet'},
        {'block_index': 1, 'text': '=Edited literal'},
        {'block_index': 2, 'text': 'Edited answer'},
    ]}
    reconstruct_from_sidecar(sidecar, {}, str(workbook), str(output))
    assert_formulas(output)
    wb = load_workbook(output)
    assert wb['Lesson']['C1'].value == 'Edited worksheet'
    assert wb['Lesson']['D1'].data_type == 's'
    wb.close()


def test_shared_formula_ranges_survive_translation(tmp_path):
    source = tmp_path / 'shared.xlsx'
    wb = Workbook()
    ws = wb.active
    ws['A1'] = '=B1*2'
    ws['A2'] = '=B2*2'
    ws['B1'] = 2
    ws['B2'] = 3
    ws['C1'] = 'Worksheet'
    wb.save(source)
    namespace = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with ZipFile(source) as z:
        files = {name: z.read(name) for name in z.namelist()}
    xml = ET.fromstring(files['xl/worksheets/sheet1.xml'])
    for address in ['A1', 'A2']:
        formula = xml.find(f'.//s:c[@r="{address}"]/s:f', namespace)
        formula.set('t', 'shared')
        formula.set('si', '0')
        if address == 'A1': formula.set('ref', 'A1:A2')
        else: formula.text = None
    files['xl/worksheets/sheet1.xml'] = ET.tostring(xml)
    with ZipFile(source, 'w', ZIP_DEFLATED) as z:
        for name, value in files.items(): z.writestr(name, value)
    output = tmp_path / 'translated.xlsx'
    translate_xlsx_inplace(source, output, lambda text, **kw: 'Mga buluhaton')
    wb = load_workbook(output)
    assert wb.active['A1'].value == '=B1*2'
    assert wb.active['A2'].value == '=B2*2'
    assert wb.active['C1'].value == 'Mga buluhaton'
    wb.close()


def test_formula_only_workbook_never_calls_translation(tmp_path):
    source, output = tmp_path / 'formula.xlsx', tmp_path / 'out.xlsx'
    wb = Workbook()
    wb.active['A1'] = '=1+1'
    wb.save(source)
    def forbidden(*args, **kwargs):
        pytest.fail('Formula sent to translation')
    translate_xlsx_inplace(source, output, forbidden)
    wb = load_workbook(output)
    assert wb.active['A1'].value == '=1+1'
    wb.close()


def test_new_formula_values_are_not_fabricated_in_cache(workbook, tmp_path):
    output = tmp_path / 'out.xlsx'
    translate_xlsx_inplace(workbook, output, lambda text, **kw: text)
    assert_formulas(output)
    wb = load_workbook(output, data_only=True)
    assert wb['Lesson']['A1'].value is None
    wb.close()

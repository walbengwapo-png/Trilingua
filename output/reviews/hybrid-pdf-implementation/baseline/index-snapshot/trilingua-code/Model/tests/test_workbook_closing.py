"""Offline workbook calculation checks; never calls a provider."""
import os
from pathlib import Path
import shutil
import subprocess
import sys
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.formula import ArrayFormula

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from document.reconstructor import translate_xlsx_inplace, write_xlsx


@pytest.mark.parametrize('writer', ['translation', 'legacy-original', 'legacy-new'])
def test_every_writer_requests_automatic_full_recalculation(tmp_path, writer):
    source, output = tmp_path / 'original.xlsx', tmp_path / 'translated.xlsx'
    wb = Workbook()
    wb.active['A1'] = '=1+1'
    wb.active['B1'] = 'Worksheet'
    wb.calculation.calcMode = 'manual'
    wb.calculation.fullCalcOnLoad = False
    wb.calculation.forceFullCalc = False
    wb.save(source)
    if writer == 'translation':
        translate_xlsx_inplace(source, output, lambda text, **kw: 'Mga buluhaton')
    else:
        write_xlsx([], output, source if writer == 'legacy-original' else None)
    wb = load_workbook(output)
    assert wb.calculation.calcMode == 'auto'
    assert wb.calculation.fullCalcOnLoad is True
    assert wb.calculation.forceFullCalc is True
    wb.close()


def test_libreoffice_recalculated_standard_cross_sheet_array_and_shared_values(tmp_path):
    executable = os.environ.get('TRILINGUA_LIBREOFFICE') or shutil.which('soffice')
    if not executable:
        candidate = Path(r'C:\Program Files\LibreOffice\program\soffice.com')
        executable = str(candidate) if candidate.exists() else None
    if not executable:
        pytest.skip('LibreOffice absent: original/translated computed-value comparison cannot run; install nothing.')
    source, translated = tmp_path / 'original.xlsx', tmp_path / 'translated.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.title = 'Lesson'
    ws['A1'], ws['A2'] = 2, 3
    ws['B1'] = '=SUM(A1:A2)'
    ws['C1'] = 'Worksheet'
    ws['D1'] = ArrayFormula(ref='D1:D2', text='=A1:A2*2')
    ws['E1'], ws['E2'] = '=A1*2', '=A2*2'
    wb.create_sheet('Answers')['A1'] = '=Lesson!B1'
    # LibreOffice conversion respects manual mode rather than recalculating it.
    # The separate writer tests above exercise manual-mode inputs.
    wb.calculation.calcMode = 'auto'
    wb.calculation.fullCalcOnLoad = True
    wb.save(source)
    with ZipFile(source) as z:
        entries = {n: z.read(n) for n in z.namelist()}
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    xml = ET.fromstring(entries['xl/worksheets/sheet1.xml'])
    for address in ['E1', 'E2']:
        formula = xml.find(f'.//s:c[@r="{address}"]/s:f', ns)
        formula.set('t', 'shared')
        formula.set('si', '0')
        if address == 'E1': formula.set('ref', 'E1:E2')
        else: formula.text = None
    entries['xl/worksheets/sheet1.xml'] = ET.tostring(xml)
    with ZipFile(source, 'w', ZIP_DEFLATED) as z:
        for name, value in entries.items(): z.writestr(name, value)
    translate_xlsx_inplace(source, translated, lambda text, **kw: 'Mga buluhaton')
    values = []
    for i, path in enumerate([source, translated]):
        destination = tmp_path / f'calculated-{i}'
        destination.mkdir()
        profile = (tmp_path / f'lo-profile-{i}').as_uri()
        result = subprocess.run([executable, f'-env:UserInstallation={profile}', '--headless',
            '--convert-to', 'xlsx', '--outdir', str(destination), str(path)],
            capture_output=True, text=True, timeout=60)
        assert result.returncode == 0, result.stdout + result.stderr
        calculated = destination / path.name
        assert calculated.exists(), result.stdout + result.stderr
        book = load_workbook(calculated, data_only=True)
        values.append([book['Lesson'][c].value for c in ['B1', 'D1', 'D2', 'E1', 'E2']] + [book['Answers']['A1'].value])
        book.close()
    assert values[0] == [5, 4, 6, 4, 6, 5]
    assert values[1] == values[0]

# -*- coding: utf-8 -*-
"""
Tests for the parallel in-place DOCX translation path.

The in-place path (DOCX/PPTX/XLSX) used to call the LLM once per text element,
sequentially. It now runs collect -> parallel translate (batch_translate_blocks)
-> replay. These tests verify:

1. The parallel path is actually used (batch_translate_blocks is invoked).
2. Round-trip correctness: header, paragraphs, and table cells all get the
   translated text, in the same order the admin sidecar expects.
3. The collect pass (save=False) does not write an output file.
4. The prepass context preamble is threaded into worker translations.
5. Run-level formatting (bold) survives the replay pass.

Run:
    python -m pytest trilingua-code/Model/tests/test_docx_inplace_parallel.py -v
"""

import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import pytest
from unittest import mock
from docx import Document

from dto.requests import DocumentTranslationRequest
from document.reconstructor import _translate_docx_inplace_with_translator


def _mock_translation(text: str) -> str:
    """Mirror MockTranslationProvider output for expected-value checks."""
    return f"[MOCK:{len(text.split())}w] {text[::-1]}"


def _make_docx(path, header="Header Text",
               paragraphs=("Paragraph One text.", "Second paragraph here."),
               table_data=(("Alpha Cell", "Beta Cell"),
                           ("Gamma Cell", "Delta Cell"))):
    """Build a DOCX with a header, body paragraphs, and a 2x2 table."""
    doc = Document()
    section = doc.sections[0]
    section.header.is_linked_to_previous = False
    section.header.paragraphs[0].text = header
    for text in paragraphs:
        doc.add_paragraph(text)
    tbl = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
    for r, row_data in enumerate(table_data):
        for c, cell_text in enumerate(row_data):
            tbl.rows[r].cells[c].text = cell_text
    doc.save(path)
    return path


def _docx_texts(path):
    """Return (header_text, body_paragraph_texts, table_cell_texts) from a DOCX."""
    doc = Document(path)
    header = doc.sections[0].header.paragraphs[0].text
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    return header, paras, cells


# ===========================================================================
# Parallel path is used
# ===========================================================================

def test_inplace_uses_batch_translate_blocks(document_pipeline, tmp_path):
    """The in-place path must route through batch_translate_blocks."""
    src = _make_docx(str(tmp_path / "in.docx"))
    request = DocumentTranslationRequest(
        file_path=src, source_lang="English", target_lang="Cebuano",
    )

    with mock.patch.object(
        document_pipeline.translation_pipeline, "batch_translate_blocks",
        wraps=document_pipeline.translation_pipeline.batch_translate_blocks,
    ) as spy:
        response = document_pipeline.translate(request)

    assert response.success, response.error_message
    spy.assert_called_once()
    assert response.sidecar is not None
    assert os.path.exists(response.output_path)


# ===========================================================================
# Round-trip content + sidecar ordering
# ===========================================================================

def test_inplace_docx_roundtrip_content(document_pipeline, tmp_path):
    """Header, paragraphs, and table cells all receive translated text."""
    src = _make_docx(str(tmp_path / "in.docx"))
    request = DocumentTranslationRequest(
        file_path=src, source_lang="English", target_lang="Cebuano",
    )
    response = document_pipeline.translate(request)

    header, paras, cells = _docx_texts(response.output_path)
    assert header == _mock_translation("Header Text")
    assert paras == [_mock_translation(p) for p in ("Paragraph One text.", "Second paragraph here.")]
    assert cells == [_mock_translation(c) for c in
                     ("Alpha Cell", "Beta Cell", "Gamma Cell", "Delta Cell")]


def test_inplace_sidecar_block_order(document_pipeline, tmp_path):
    """block_index is sequential and source order is header -> paragraphs -> cells."""
    src = _make_docx(str(tmp_path / "in.docx"))
    request = DocumentTranslationRequest(
        file_path=src, source_lang="English", target_lang="Cebuano",
    )
    response = document_pipeline.translate(request)

    blocks = response.sidecar["blocks"]
    expected_sources = (
        ["Header Text"]
        + ["Paragraph One text.", "Second paragraph here."]
        + ["Alpha Cell", "Beta Cell", "Gamma Cell", "Delta Cell"]
    )
    assert [b["block_index"] for b in blocks] == list(range(len(expected_sources)))
    assert [b["source_text"] for b in blocks] == expected_sources
    assert blocks[0]["block_type"] == "header"
    assert blocks[-1]["block_type"] == "table_cell"


# ===========================================================================
# Collect pass does not save
# ===========================================================================

def test_collect_pass_does_not_save(tmp_path):
    """save=False walks the document and collects text without writing output."""
    src = _make_docx(str(tmp_path / "in.docx"))
    out = str(tmp_path / "out.docx")

    collected = []

    def recorder(text, block_type="paragraph"):
        collected.append((block_type, text))
        return text

    _translate_docx_inplace_with_translator(src, out, recorder, save=False)
    assert not os.path.exists(out), "collect pass must not write output"
    assert collected[0][0] == "header"
    assert [t for _, t in collected] == (
        ["Header Text"]
        + ["Paragraph One text.", "Second paragraph here."]
        + ["Alpha Cell", "Beta Cell", "Gamma Cell", "Delta Cell"]
    )

    # save=True (default) still writes the file
    _translate_docx_inplace_with_translator(src, out, recorder, save=True)
    assert os.path.exists(out)


# ===========================================================================
# Prepass context preamble reaches workers
# ===========================================================================

def test_context_preamble_injected_into_worker(translation_pipeline):
    """_translate_single_worker prepends the preamble to the context hint."""
    provider = translation_pipeline.provider
    captured = {}
    with mock.patch.object(provider, "translate", wraps=provider.translate) as spy:
        translation_pipeline._translate_single_worker(
            "Hello world", "English", "Cebuano", "paragraph",
            document_memory=None, translation_cache=None,
            quality_reviewer=None, document_profile=None,
            block_index=0, context_preamble="[PREPASS] mock summary",
        )
        _, kwargs = spy.call_args
    assert "[PREPASS]" in kwargs.get("context_hint", "")


def test_context_preamble_propagates_through_batch(translation_pipeline):
    """batch_translate_blocks hands the preamble to single-item workers."""
    blocks = [{"type": "paragraph", "text": "A reasonably long block that "
                                            "must travel alone because it "
                                            "exceeds the batch budget so it "
                                            "becomes its own single batch item "
                                            "for this unit test of preamble "
                                            "propagation through the executor "
                                            "infrastructure."}]
    with mock.patch.object(
        translation_pipeline, "_translate_single_worker",
        wraps=translation_pipeline._translate_single_worker,
    ) as spy:
        out = translation_pipeline.batch_translate_blocks(
            blocks, "English", "Cebuano", context_preamble="[PREPASS] terms",
        )
    assert len(out) == 1
    args = spy.call_args.args
    assert args[9] == "[PREPASS] terms", "context_preamble passed positionally to worker"


# ===========================================================================
# Run-level formatting preserved
# ===========================================================================

def test_inplace_preserves_bold_run_formatting(document_pipeline, tmp_path):
    """Bold formatting on a run survives the replay pass."""
    src = str(tmp_path / "in.docx")
    doc = Document()
    p = doc.add_paragraph()
    r1 = p.add_run("First part is ")
    r1.bold = True
    r2 = p.add_run("second part is plain")
    doc.save(src)

    request = DocumentTranslationRequest(
        file_path=src, source_lang="English", target_lang="Cebuano",
    )
    response = document_pipeline.translate(request)

    out = Document(response.output_path)
    para = out.paragraphs[0]
    assert len(para.runs) == 2, "run count must be preserved"
    assert para.runs[0].bold is True

import sys
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from dto.responses import TranslationResponse
from dto.requests import TranslationRequest
from pipeline.document_context import DocumentContext
from pipeline.document_pipeline import _engine_metrics
from pipeline.translation_pipeline import TranslationPipeline
from providers.gptoss import GPTOSSProvider
from providers.gemini import GeminiProvider


def test_provider_models_are_not_overwritten():
    ctx=DocumentContext()
    ctx.record_provider('fixture','one')
    ctx.record_provider('fixture','two')
    ctx.record_provider('fallback','three')
    assert ctx.providers_used['fixture']['model']=='one'
    assert ctx.summary()['observed_providers']==[
        {'provider':'fixture','model':'one'}, {'provider':'fixture','model':'two'},
        {'provider':'fallback','model':'three'}]


def test_document_envelope_preserves_attribution_without_claiming_http_counts():
    ctx=DocumentContext()
    ctx.record_provider('fixture','one')
    metadata=_engine_metrics(ctx)['measurement_metadata']
    assert metadata['schema_version']==1
    assert metadata['observed_providers']==[{'provider':'fixture','model':'one'}]
    assert metadata['calls_provenance']=='unknown'
    assert metadata['calls_complete'] is False
    assert metadata['timing_scope']=='document_pipeline_elapsed'


def test_cache_marks_observed_zero_translation_calls(monkeypatch):
    provider=SimpleNamespace(name='fixture',model_name='one',estimate_tokens=lambda s:1)
    pipeline=TranslationPipeline(provider)
    monkeypatch.setattr(pipeline,'_get_cache',lambda:SimpleNamespace(get=lambda *a,**kw:'Kumusta mga estudyante.'))
    result=pipeline.translate(TranslationRequest(text='Hello students.',source_lang='English',target_lang='Cebuano'))
    metadata=result.measurement_metadata
    assert metadata['cache_served'] is True
    assert metadata['llm_calls']==0 and metadata['calls_complete'] is True
    assert metadata['calls_scope']=='translation'
    assert metadata['observed_providers']==[]
    assert result.token_usage=={}


@pytest.mark.parametrize('usage,expected',[
    ({'prompt_eval_count':12,'eval_count':5},{'input':'provider-reported','output':'provider-reported'}),
    ({'prompt_eval_count':12},{'input':'provider-reported','output':'unknown'}),
    ({'eval_count':5},{'input':'unknown','output':'provider-reported'}),
    ({},{'input':'estimated','output':'estimated'}),
])
def test_gptoss_usage_provenance_is_per_direction(usage,expected):
    provider=GPTOSSProvider(api_url='http://127.0.0.1:19999/fixture',model='fixture-model')
    provider._session.post=Mock(return_value=SimpleNamespace(status_code=200,headers={},raise_for_status=lambda:None,
        json=lambda:{'message':{'content':'Kumusta'},**usage}))
    result=provider.translate('Hello','English','Cebuano')
    assert result.success
    assert provider._session.post.call_count==1
    assert result.measurement_metadata['usage_provenance']==expected


def test_gemini_current_proxies_are_labelled_estimated():
    provider=GeminiProvider(api_key='fixture-key',model='fixture-model')
    provider._session.post=Mock(return_value=SimpleNamespace(status_code=200,headers={},raise_for_status=lambda:None,
        json=lambda:{'candidates':[{'content':{'parts':[{'text':'Kumusta'}]},'finishReason':'STOP'}]}))
    result=provider.translate('Hello','English','Cebuano')
    assert result.success
    assert result.measurement_metadata['usage_provenance']=={'input':'estimated','output':'estimated'}
    assert result.measurement_metadata['usage_method']['input']=='characters-divided-by-four-v1'


def test_dto_round_trip_keeps_measurement_metadata():
    metadata={'schema_version':1,'usage_provenance':{'input':'unknown'}}
    result=TranslationResponse(translated_text='Output',measurement_metadata=metadata)
    assert result.to_dict()['measurement_metadata']==metadata


def test_cached_text_review_does_not_claim_zero_operation_calls(monkeypatch):
    import server
    cached=TranslationResponse(translated_text='Kumusta mga estudyante.',provider='cache',model='configured',
        measurement_metadata={'schema_version':1,'cache_served':True,'llm_calls':0,
            'calls_complete':True,'calls_scope':'translation','calls_provenance':'measured locally'})
    monkeypatch.setattr(server,'_translation_pipeline',SimpleNamespace(translate=lambda req:cached))
    monkeypatch.setattr(server,'_run_pipeline_guarded',lambda f:f())
    monkeypatch.setattr(server,'_text_quality_reviewer',SimpleNamespace(review=Mock(side_effect=RuntimeError('unavailable'))))
    result=server.translate_text(server.TextRequest(text='Hello students.',source_lang='English',target_lang='Cebuano',mode='balanced'))
    assert result['measurement_metadata']['calls_complete'] is False
    assert result['measurement_metadata']['llm_calls'] is None
    assert result['measurement_metadata']['calls_scope']=='operation'


@pytest.mark.parametrize('extension',['docx','xlsx','pptx','csv','pdf'])
def test_document_format_paths_keep_phase_a_metadata(extension,tmp_path,document_pipeline,monkeypatch):
    from dto.requests import DocumentTranslationRequest
    text='The children read a short lesson.'
    source=tmp_path/('source.'+extension)
    if extension=='docx':
        from docx import Document
        document=Document(); document.add_paragraph(text); document.save(source)
    elif extension=='xlsx':
        from openpyxl import Workbook
        workbook=Workbook(); workbook.active['A1']=text; workbook.save(source); workbook.close()
    elif extension=='pptx':
        from pptx import Presentation
        from pptx.util import Inches
        presentation=Presentation(); slide=presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1),Inches(1),Inches(7),Inches(2)).text=text
        presentation.save(source)
    elif extension=='csv':
        source.write_text(text+'\n',encoding='utf8')
    else:
        import fitz
        with fitz.open() as document:
            document.new_page().insert_text((50,70),text)
            document.save(source)
    monkeypatch.setattr(document_pipeline,'_get_cache',lambda:None)
    monkeypatch.setattr(document_pipeline.translation_pipeline,'_get_cache',lambda:None)
    response=document_pipeline.translate(DocumentTranslationRequest(str(source),'English','Cebuano',mode='fast'))
    assert response.success and Path(response.output_path).is_file()
    metadata=response.to_dict()['metrics']['measurement_metadata']
    assert metadata['schema_version']==1 and metadata['calls_complete'] is False
    assert metadata['usage_provenance']=={'input':'unknown','output':'unknown'}
    assert metadata['timing_scope']=='document_pipeline_elapsed'
    assert {'provider':'mock_translation','model':'mock_model_v0'} in metadata['observed_providers']

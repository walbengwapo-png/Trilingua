"""Regression tests for the in-place PowerPoint translation walker."""

import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from pptx import Presentation
from pptx.util import Inches

from document.reconstructor import translate_pptx_inplace


def test_inplace_pptx_translates_regular_text_box(tmp_path):
    """A non-placeholder text box must not access placeholder_format."""
    source = tmp_path / "source.pptx"
    output = tmp_path / "translated.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    text_box.text_frame.paragraphs[0].text = "Regular text box"
    presentation.save(source)

    calls = []

    def translate(text, block_type="paragraph"):
        calls.append((text, block_type))
        return f"Translated: {text}"

    translate_pptx_inplace(str(source), str(output), translate)

    translated = Presentation(output)
    assert calls == [("Regular text box", "paragraph")]
    assert translated.slides[0].shapes[0].text == "Translated: Regular text box"
